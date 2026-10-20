"""Validate consequential arithmetic, document structure, and the exact excerpt."""
from pathlib import Path
import json
from pptx import Presentation
from pypdf import PdfReader
from lxml import etree
from model import compute

root=Path(__file__).parent
result=compute()
stored=json.loads((root/'model_results.json').read_text())
assert result==stored,'Regenerate model outputs after editing inputs'
for name,n in [('main',6),('banner',4),('research',5)]:
    p=Presentation(root/f'cbp-{name}-v3.pptx')
    pdf=PdfReader(root/f'cbp-{name}-v3.pdf')
    assert len(p.slides)==len(pdf.pages)==n
    for i,s in enumerate(p.slides):
        assert s.notes_slide.notes_text_frame.text.strip()
        for sh in s.shapes:
            assert sh.left>=0 and sh.top>=0
            assert sh.left+sh.width<=p.slide_width+100 and sh.top+sh.height<=p.slide_height+100,(name,i,sh.name)
            for effect in sh._element.xpath('./p:style/a:effectRef'):
                assert effect.get('idx')=='0'
    print(name,n,'pages; notes, boundaries, flat graphics passed')
p=Presentation(root/'cbp-main-v3.pptx');b=Presentation(root/'cbp-banner-v3.pptx')
for source,extracted in zip(list(p.slides)[1:5],b.slides):
    assert len(source.shapes)==len(extracted.shapes)
    assert all(etree.tostring(x.element)==etree.tostring(y.element) for x,y in zip(source.shapes,extracted.shapes))
    assert source.notes_slide.notes_text_frame.text==extracted.notes_slide.notes_text_frame.text
print('Exact excerpt: passed; illustrative numerical model: passed')
