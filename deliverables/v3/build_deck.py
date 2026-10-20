"""CBP v3: original budget-adaptive repositioning idea, editable slide graphics."""
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from model import compute, MODES

OUT=Path(__file__).parent; DATA=compute()
FONT='Noto Sans CJK KR'
NAVY='162D4D'; INK='20354D'; TEAL='00877F'; BLUE='598C9F'; GOLD='DAA64D'
WHITE='FFFFFF'; BG='F6F8FB'; PALE='E4F2EF'; LIGHT='E8EEF5'; GRAY='63778C'; LINE='D5DFE9'; RED='B35C3A'; MUTED='AEC3D6'
COLORS=[NAVY,TEAL,GOLD]
def prs():
    p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9)
    p.core_properties.author='두런두런';p.core_properties.title='CBP | 탄소예산 기반 유닛로드시스템 운영 최적화 프레임워크'
    p.core_properties.subject='탄소예산 상태에 적응하는 파렛트 재배치 의사결정 모형'
    return p
def box(s,x,y,w,h,c=WHITE,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c)
    if line:sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    sh._element.spPr.append(OxmlElement('a:effectLst'));return sh
def tx(s,x,y,w,h,t,sz=20,c=INK,b=False,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,t0 in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t0;p.font.name=FONT;p.font.size=Pt(sz);p.font.bold=b;p.font.color.rgb=RGBColor.from_string(c)
        p.space_after=Pt(2);p.line_spacing=1.12
        if align is not None:p.alignment=align
    return sh
def line(s,x1,y1,x2,y2,c=LINE,width=1.5,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(c);sh.line.width=Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh
def circle(s,x,y,d,c):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();return sh
def arrow(s,x,y,w,c=TEAL,down=False):
    sh=s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW if down else MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(.23 if down else w),Inches(w if down else .23));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();return sh
def label(s,x,y,w,t,c=PALE,ink=TEAL,sz=12):
    box(s,x,y,w,.36,c);tx(s,x+.10,y+.055,w-.2,.25,t,sz,ink,True)
def node(s,x,y,letter,title,sub,c=TEAL):
    circle(s,x,y,.52,c);tx(s,x,y+.09,.52,.3,letter,16,WHITE,True,PP_ALIGN.CENTER)
    tx(s,x+.72,y-.02,3.25,.39,title,21,INK,True);tx(s,x+.73,y+.52,3.3,.65,sub,16,GRAY)
def warehouse(s,x,y,w=1,c=NAVY):
    box(s,x,y+.22*w,1.50*w,.91*w,c)
    sh=s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(x-.06*w),Inches(y-.12*w),Inches(1.62*w),Inches(.60*w));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background()
    box(s,x+.43*w,y+.60*w,.64*w,.53*w,'D1E0EB')
def pallet(s,x,y,w=1):
    for i in range(3):
        box(s,x,y-i*.14*w,1.05*w,.065*w,GOLD)
        for a in [.05,.47,.89]:box(s,x+a*w,y+.065*w-i*.14*w,.10*w,.07*w,GOLD)
def base(p,n,title,sub='',dark=False,appendix=False):
    s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY if dark else BG)
    tx(s,.65,.31,12,.28,'CBP  /  CARBON BUDGET POOLING',12,MUTED if dark else TEAL,True)
    tx(s,13.45,.31,1.9,.28,('연구 보충 ' if appendix else '')+f'{n:02}',12,MUTED if dark else GRAY,True,PP_ALIGN.RIGHT)
    tx(s,.65,.93,14.7,1.16,title,34,WHITE if dark else INK,True)
    if sub:tx(s,.68,2.34 if '\n' in title else 2.10,14.6,.52,sub,17,MUTED if dark else GRAY)
    line(s,.65,8.50,15.35,8.50,'3B5571' if dark else LINE,1)
    tx(s,.67,8.64,14.6,.21,'두런두런  |  제13회 유닛로드시스템(ULS) 창의성 공모전  |  탄소예산 기반 재배치 의사결정',10,MUTED if dark else GRAY)
    return s
def foot(s,t):tx(s,.73,8.17,14.55,.24,t,10,GRAY)
def note(s,t):s.notes_slide.notes_text_frame.text=t
def save_clean(p,path):
    for s in p.slides:
        for sh in s.shapes:
            for ref in sh._element.xpath('./p:style/a:effectRef'):
                ref.set('idx','0')
    p.save(path)

SCRIPTS=[
('00:00–00:30','문제와 연구 질문','같은 파렛트 천 매를 옮기더라도, 이번 달 탄소예산을 얼마나 썼는지에 따라 선택이 달라져야 하지 않을까요? 저희 CBP는 탄소예산을 운영의 신호로 쓰는 아이디어입니다. 예산이 충분할 때는 비용에, 많이 썼을 때는 탄소에 더 무게를 두어, 파렛트 재배치의 우선순위를 조정합니다.'),
('00:30–01:10','닫힌 고리로 작동하는 CBP','작동 방식은 이 순환입니다. 운영자가 기간별 탄소예산을 정하고, 지금까지 쓴 배출량을 나누면 예산 사용률이 나옵니다. 예산 만 킬로그램 중 팔천이백을 썼다면 팔십이 퍼센트입니다. 이 상태에 맞춰 운영 모드를 정하고 운송안을 비교합니다. 단, 필요한 수량과 도착 기한, 적재 조건을 만족하는 안만 비교합니다. 실행 뒤에는 실제 배출을 반영해 다음 선택을 바꿉니다.'),
('01:10–01:45','네 가지 모드','원안의 네 가지 모드입니다. 사용률이 올라갈수록 비용의 비중은 낮추고 탄소의 비중은 높입니다. 예를 들어 탄소 모드에서는 비용 이십, 탄소 육십, 도착 시간의 여유 이십 퍼센트를 반영합니다. 이 경계값과 가중치는 초기 설계값이며, 현장 자료로 조정해야 합니다. 어떤 모드에서도 필요한 수량과 기한을 자동으로 포기하지 않습니다.'),
('01:45–02:35','재현 가능한 숫자 예시','천 매를 여덟 시간 안에 보내는 세 가지 운송안이 있다고 가정하겠습니다. A는 백만 원에 배출 사백, B는 백십만 원에 이백육십, C는 백사십만 원에 이백입니다. 각 항목을 영에서 일 사이 점수로 맞추고, 모드별 비중을 곱해 더합니다. 현재 사용률 팔십이 퍼센트의 탄소 모드에서는 C가 영 점 칠이오로 가장 높습니다. 계산 방법과 입력값을 공개해 같은 결과를 다시 확인할 수 있게 했습니다.'),
('02:35–03:15','같은 후보, 달라지는 선택','같은 후보를 놓고 비교하면 비용 모드에서는 A, 균형 모드에서는 B, 탄소 모드에서는 C를 고릅니다. 핵심은 무조건 가장 싼 안도, 무조건 배출이 적은 안도 고정하지 않는다는 점입니다. 물론 대가가 있습니다. 이 예시에서 A 대신 C를 고르면 비용 사십만 원이 더 들고 배출 이백 킬로그램이 줄어듭니다. 이는 실측 성과가 아니라, 선택의 원리를 보여 주는 계산 예시입니다.'),
('03:15–04:00','실증 계획과 기여','현장에서는 동일 풀링사의 호환 파렛트부터 시작합니다. 먼저 실제 배차에 영향을 주지 않는 추천으로 시험하고, 비용 중심 방식과 고정 가중치 방식, CBP를 같은 요청과 후보로 비교하겠습니다. 평가 기준은 총비용, 실제 배출, 정시 공급, 기간말 예산 초과입니다. 남은 기간과 수요도 함께 검토해야 합니다. CBP의 제안은 명확합니다. 탄소를 보고서에 남는 숫자에서, 다음 파렛트 이동을 바꾸는 의사결정 신호로 만드는 것입니다.')]

p=prs()
# 1 — A research question grounded in a concrete ULS movement.
s=base(p,1,'같은 파렛트 요청,\n탄소예산이 다르면 선택도 달라져야 합니다.',dark=True)
tx(s,.72,2.55,6.25,.46,'탄소예산 기반 유닛로드시스템\n운영 최적화 프레임워크',21,WHITE,True)
tx(s,.73,3.83,5.8,1.08,'탄소예산 사용률에 따라\n재배치의 우선순위를 바꾸는 CBP',27,'8BDFCC',True)
tx(s,.74,5.30,5.75,1.28,'예산 여유가 크면 → 비용에 무게\n예산을 많이 썼으면 → 탄소에 무게',21,WHITE)
tx(s,.75,7.05,5.55,.69,'연구 질문\n“탄소 상태를 다음 운송 선택에 어떻게 연결할까?”',16,MUTED)
box(s,7.30,2.65,8.02,4.98,'233E5D')
tx(s,7.68,2.97,7.26,.35,'1,000매를 보내는 세 가지 운송안',22,WHITE,True)
for i,(a,yy,desc,c) in enumerate([('A',3.83,'저비용형',MUTED),('B',4.97,'절충형','8BDFCC'),('C',6.11,'저탄소형',GOLD)]):
    warehouse(s,7.80,yy,.48,c);tx(s,8.82,yy+.1,2,.3,a+' 거점',17,WHITE,True)
    line(s,10.20,yy+.30,12.41,5.25,c,2);tx(s,10.10,yy-.13,2.15,.3,desc,15,c)
warehouse(s,13.05,4.66,.85,'8BDFCC');pallet(s,13.29,6.15,.70)
tx(s,12.74,6.54,2.19,.35,'수요 거점',18,WHITE,True,PP_ALIGN.CENTER)
tx(s,7.70,7.16,7.22,.25,'개념도 · 운송안은 재고·노선·차량 조합으로 구성',12,MUTED)
tx(s,7.38,7.92,7.89,.28,'예선 원안 계승  /  CUR · 4단계 운영 모드 · RPS',14,'8BDFCC',True)

# 2 — State → policy → action → feedback; hard feasibility stays separate.
s=base(p,2,'탄소예산을 “운영 신호”로 바꾸는 순환',
       'CBP는 배출 실적을 다음 파렛트 재배치의 의사결정에 다시 연결합니다.')
box(s,.73,2.89,3.60,4.86,NAVY)
tx(s,1.03,3.18,2.98,.31,'01  현재 상태 읽기',18,MUTED,True)
tx(s,1.00,3.77,3.00,.91,'82%',61,WHITE,True)
tx(s,1.04,4.88,2.96,.36,'탄소예산 사용률 CUR',17,'8BDFCC',True)
box(s,1.05,5.53,2.94,.16,'486079');box(s,1.05,5.53,2.94*.82,.16,'8BDFCC')
tx(s,1.05,6.05,2.95,.89,'사용 8,200 ÷ 예산 10,000\n× 100 = 82%',17,WHITE)
tx(s,1.05,7.22,2.95,.26,'단위 kgCO₂e · 설명용 가정',11,MUTED)
arrow(s,4.57,4.08,.62)
for x,n,title,body in [(5.43,'02','모드 결정','사용률 구간에 맞춰\n비용·탄소 비중 조정'),(10.62,'03','후보 평가','제약을 통과한 후보만\n종합점수를 비교해 선택')]:
    box(s,x,3.20,4.65,2.30,WHITE,LINE);label(s,x+.24,3.43,.57,n,NAVY,WHITE)
    tx(s,x+.98,3.43,3.42,.42,title,25,INK,True);tx(s,x+.26,4.20,4.13,.83,body,20,GRAY)
arrow(s,10.19,4.08,.29)
box(s,5.43,5.83,9.84,.70,PALE)
tx(s,5.68,6.02,9.34,.33,'후보 진입 조건  |  수량 · 기한 · 호환 규격 · 적재 · 잔여예산',17,TEAL,True)
line(s,12.94,5.51,12.94,5.80,TEAL,2)
box(s,5.43,6.95,9.84,.82,NAVY)
tx(s,5.70,7.17,9.30,.37,'04  실행 → 실제 배출 반영 → 다음 의사결정',23,WHITE,True)
line(s,5.43,7.38,4.73,7.38,TEAL,2);line(s,4.73,7.38,4.73,6.53,TEAL,2);line(s,4.73,6.53,4.33,6.53,TEAL,2)
sh=arrow(s,4.33,6.41,.40);sh.rotation=180
arrow(s,12.82,6.59,.27,down=True)
foot(s,'예산은 운영자가 정하는 내부 관리 한도입니다. 법정 배출권을 뜻하지 않으며, 현재 단계는 가중치 기반 의사결정 모형입니다.')

# 3 — Original four modes, visible proportions, thresholds treated as hypotheses.
s=base(p,3,'예산을 많이 쓸수록, 탄소의 비중을 높입니다.',
       '재배치 우선순위 점수(RPS): 세 기준의 가중합. 한도에 닿기 전부터 탄소 비중을 높입니다.')
for xx,c,t in [(9.10,NAVY,'비용'),(10.89,TEAL,'탄소'),(12.69,GOLD,'시간 여유')]:
    box(s,xx,2.77,.19,.19,c);tx(s,xx+.29,2.72,1.85,.3,t,14,GRAY)
names=['비용 모드','균형 모드','탄소 모드','긴급 모드'];ranges=['CUR < 60%','60% ≤ CUR < 75%','75% ≤ CUR < 90%','CUR ≥ 90%']
descs=['운영 비용에 무게','비용과 탄소 절충','탄소 절감에 무게','탄소 비중을 더 높임']
for i,((name,weights),title,rng,desc) in enumerate(zip(MODES,names,ranges,descs)):
    x=.73+i*3.69
    box(s,x,3.26,3.45,3.95,PALE if i==2 else WHITE,TEAL if i==2 else LINE)
    tx(s,x+.22,3.51,3,.29,'0'+str(i+1)+'  /  '+name.upper(),13,TEAL if i==2 else GRAY,True)
    tx(s,x+.22,4.06,3,.51,title,28,INK,True)
    tx(s,x+.22,4.78,3,.32,rng,16,TEAL,True)
    pos=x+.22
    for val,c in zip(weights,COLORS):
        box(s,pos,5.45,3.01*val,.50,c);tx(s,pos,5.56,3.01*val,.26,str(round(val*100)),11,WHITE if c!=GOLD else NAVY,True,PP_ALIGN.CENTER);pos+=3.01*val
    tx(s,x+.22,6.22,3.02,.34,desc,17,INK,True)
    if i==2:label(s,x+.22,6.74,2.98,'예시 CUR 82% → 이 모드',TEAL,WHITE,12)
box(s,.73,7.53,14.53,.48,LIGHT)
tx(s,.94,7.63,14.12,.28,'RPS = 비용점수 × 비용비중 + 탄소점수 × 탄소비중 + 시간여유점수 × 시간비중',18,INK,True)
foot(s,'구간과 가중치는 예선 원안의 초기 설계 가정입니다. 검증된 최적값이 아니며, 긴급 모드도 수량·기한을 자동 완화하지 않습니다.')

# 4 — Numerical example with all inputs and score contributions.
s=base(p,4,'CUR 82%일 때, 세 운송안 중 무엇을 고를까요?',
       '모두 1,000매를 8시간 이내 공급할 수 있다고 가정합니다. 입력값은 실측·견적이 아닌 설명용입니다.')
label(s,.74,2.82,5.88,'동일 요청 · 동일 후보 · 동일 계산 기준',LIGHT,NAVY,14)
cols=[(.93,1.05,'운송안'),(2.19,1.46,'비용\n만원'),(3.86,1.73,'배출\nkgCO₂e'),(5.80,1.28,'도착\n시간')]
for x,w,t in cols:tx(s,x,3.42,w,.65,t,15,GRAY,True,PP_ALIGN.CENTER)
for i,(a,c,e,h) in enumerate([('A',100,400,6),('B',110,260,4),('C',140,200,3)]):
    y=4.26+i*.82;box(s,.75,y,6.47,.66,PALE if i==2 else WHITE)
    for (x,w,_),v in zip(cols,[a,str(c),str(e),str(h)]):tx(s,x,y+.13,w,.38,v,22,TEAL if i==2 else INK,True,PP_ALIGN.CENTER)
tx(s,.94,7.02,6.13,.65,'비용·배출은 작을수록 높은 점수\n시간 여유는 클수록 높은 점수',18,GRAY)
line(s,7.58,2.88,7.58,7.83,LINE,1)
tx(s,8.04,2.90,7.12,.4,'탄소 모드: 비용 20 · 탄소 60 · 시간 20',20,INK,True)
for i,(a,components) in enumerate(zip('ABC',[[.2,0,.05],[.15,.42,.10],[0,.60,.125]])):
    y=3.95+i*1.12
    tx(s,8.06,y+.10,.40,.40,a,23,TEAL if i==2 else INK,True)
    box(s,8.75,y,5.06,.64,LIGHT);pos=8.75
    for val,c in zip(components,COLORS):
        if val:box(s,pos,y,5.06*val,.64,c);pos+=5.06*val
    tx(s,14.00,y+.10,1.13,.4,f'{sum(components):.3f}',23,TEAL if i==2 else INK,True)
tx(s,8.78,7.40,6.40,.39,'C: 0.2 × 0 + 0.6 × 1 + 0.2 × 0.625 = 0.725',16,TEAL,True)
for xx,c,t in [(8.83,NAVY,'비용'),(10.45,TEAL,'탄소'),(12.08,GOLD,'시간 여유')]:box(s,xx,7.00,.16,.16,c);tx(s,xx+.25,6.94,1.65,.28,t,12,GRAY)
foot(s,'점수 범위 0–1: 비용·배출은 후보 간 역방향 min–max, 시간 여유는 (8−도착시간)/8. 상세 산식·전체 점수는 보충자료에 제시.')

# 5 — Native editable scatter chart and mode-switch strip.
s=base(p,5,'같은 후보라도, 예산 상태가 바뀌면 선택이 달라집니다.',
       'CBP의 역할은 비용과 탄소 사이의 선택 기준을 탄소예산 상태에 맞춰 조정하는 것입니다.')
tx(s,.86,2.86,7.70,.33,'세 운송안의 비용–배출 관계',19,INK,True)
x0,y0=1.38,6.79;cw,ch=6.51,3.0
line(s,x0,y0,x0+cw,y0,GRAY,1.5);line(s,x0,y0,x0,3.50,GRAY,1.5)
for e in (200,300,400):
    yy=y0-(e-150)/300*ch;line(s,x0,yy,x0+cw,yy,LINE,1,True);tx(s,.70,yy-.14,.55,.3,str(e),12,GRAY,False,PP_ALIGN.RIGHT)
for c in (100,110,120,130,140):
    xx=x0+(c-95)/50*cw;tx(s,xx-.28,6.97,.56,.27,str(c),12,GRAY,False,PP_ALIGN.CENTER)
tx(s,1.39,7.41,6.51,.31,'비용 (만원)  →',14,GRAY,False,PP_ALIGN.RIGHT)
tx(s,.90,3.28,4.37,.25,'배출 (kgCO₂e)',12,GRAY)
points=[]
for a,c,e,color in [('A',100,400,NAVY),('B',110,260,BLUE),('C',140,200,TEAL)]:
    xx=x0+(c-95)/50*cw;yy=y0-(e-150)/300*ch;points.append((xx,yy))
    circle(s,xx-.095,yy-.095,.19,color);tx(s,xx+.16,yy-.23,.55,.37,a,22,color,True)
for (x1,y1),(x2,y2) in zip(points,points[1:]):line(s,x1+.12,y1+.08,x2-.12,y2-.08,LINE,1.5,True)
tx(s,3.46,3.88,4.20,.62,'배출이 적은 선택에는\n추가 비용이 따를 수 있습니다.',17,GRAY)
line(s,8.52,2.91,8.52,7.91,LINE,1)
tx(s,9.04,2.88,6.1,.4,'4개 모드의 선택 결과',21,INK,True)
for i,(name,selected,em) in enumerate([('비용 모드','A','100만원 · 400kg'),('균형 모드','B','110만원 · 260kg'),('탄소 모드','C','140만원 · 200kg'),('긴급 모드','C','140만원 · 200kg')]):
    yy=3.54+i*.67;box(s,9.03,yy,6.20,.52,PALE if selected=='C' else WHITE)
    tx(s,9.22,yy+.10,2.2,.30,name,16,INK,True);tx(s,11.52,yy+.08,.39,.32,selected,19,TEAL,True);tx(s,12.16,yy+.13,2.87,.27,em,13,GRAY)
box(s,9.03,6.62,6.20,1.29,NAVY)
tx(s,9.28,6.82,5.69,.30,'이 가상 사례에서 A → C로 바꾸면',15,MUTED)
tx(s,9.26,7.30,5.70,.40,'비용 +40만원  /  배출 −200kgCO₂e',22,WHITE,True)
foot(s,'각 모드에서 같은 3개 후보가 모두 허용된 경우의 단일 의사결정 비교입니다. 기간 전체 감축률이나 비용 절감 효과를 입증한 결과가 아닙니다.')

# 6 — Honest practical implementation and a falsifiable evaluation.
s=base(p,6,'탄소를 보고하는 숫자에서,\n다음 파렛트 이동을 바꾸는 신호로.',
       '제안의 초점: 탄소예산 상태와 파렛트 풀링의 재배치 규칙을 연결합니다.')
for x,n,title,desc in [( .74,'01','작게 적용','동일 풀링사 · 호환 파렛트\n재고·운행·비용 자료 연결'),(5.69,'02','추천부터 검증','실제 배차에 앞서 모의 추천\n담당자가 실행 가능성 확인'),(10.64,'03','같은 조건으로 비교','비용 우선 / 고정 가중치 / CBP\n요청·후보·정보·제약 통일')]:
    box(s,x,3.00,4.61,2.38,WHITE,LINE);label(s,x+.25,3.27,.55,n,NAVY,WHITE)
    tx(s,x+.26,3.91,4.09,.42,title,25,INK,True);tx(s,x+.26,4.57,4.09,.68,desc,16,GRAY)
metrics=[('총비용','허용 가능한가?'),('실제 배출','얼마나 달라지는가?'),('정시 공급','필요량을 지켰는가?'),('기간말 예산','초과를 줄였는가?')]
for i,(a,b) in enumerate(metrics):
    x=.74+i*3.69;box(s,x,5.80,3.45,1.33,PALE);tx(s,x+.20,6.02,3.04,.35,a,21,TEAL,True);tx(s,x+.20,6.55,3.04,.30,b,15,INK)
tx(s,.80,7.58,14.4,.39,'다음 검증  |  가중치·경계값 조정, 남은 기간·예정 수요 반영, 실제 배출 오차 점검',19,INK,True)
foot(s,'현재 성숙도: 개념 모형과 재현 가능한 가상 계산. 실제 현장 효과·최적성·선행 기술 대비 독창성은 추가 검증이 필요합니다.')
for s,(_,title,script) in zip(p.slides,SCRIPTS):note(s,title+'\n\n'+script)
save_clean(p,OUT/'cbp-main-v3.pptx')

# Exact 4-page excerpt, as required by the competition notice.
b=prs()
for idx in [1,2,3,4]:
    src=p.slides[idx];dst=b.slides.add_slide(b.slide_layouts[6]);dst.background.fill.solid();dst.background.fill.fore_color.rgb=src.background.fill.fore_color.rgb
    for shape in src.shapes:dst.shapes._spTree.insert_element_before(deepcopy(shape.element),'p:extLst')
    note(dst,src.notes_slide.notes_text_frame.text)
save_clean(b,OUT/'cbp-banner-v3.pptx')

# Research appendix, kept outside the four-minute main presentation.
q=prs()
s=base(q,1,'모형 정의: 실행 가능한 안을 먼저 남기고, 점수로 고릅니다.',
       '범위: 동일 풀링사의 호환 파렛트 재배치. j는 필요한 물량을 충족하는 하나의 완결 운송안입니다.',appendix=True)
box(s,.75,2.91,6.67,4.99,WHITE,LINE)
tx(s,1.02,3.19,6.12,.39,'상태와 결정',23,INK,True)
for yy,t in [(3.95,'B : 기간별 내부 탄소예산 (kgCO₂e)'),(4.51,'Eₜ : 현재까지의 실제 누적 배출'),(5.07,'CURₜ = 100 × Eₜ / B  (B > 0)'),(5.63,'mₜ : CUR 구간에 따른 운영 모드'),(6.19,'j* = arg maxⱼ∈Fₜ RPSⱼ(mₜ)')]:tx(s,1.05,yy,6.10,.40,t,20,INK)
tx(s,1.05,7.03,6.10,.54,'Fₜ: 수량·기한·호환·적재·잔여예산을\n충족하는 후보 집합',16,GRAY)
box(s,7.82,2.91,7.43,4.99,PALE)
tx(s,8.10,3.19,6.86,.39,'후보별 0–1 점수',23,TEAL,True)
for yy,t in [(3.95,'비용: (Cmax − Cⱼ) / (Cmax − Cmin)'),(4.65,'탄소: (emax − eⱼ) / (emax − emin)'),(5.35,'시간 여유: (D − Tⱼ) / D'),(6.05,'RPSⱼ = α·비용점수 + β·탄소점수'),(6.49,'               + γ·시간여유점수')]:tx(s,8.11,yy,6.86,.44,t,20,INK)
tx(s,8.12,7.12,6.85,.50,'C: 비용 / e: 예상 배출 / T: 도착 소요시간\nD: 허용시간 > 0 / α+β+γ=1, 각 가중치 ≥ 0',14,GRAY)
foot(s,'min·max는 같은 실행 가능 후보 집합에서 산출. 전부 같은 값이면 해당 항목은 모두 1점. 동점은 배출 → 비용 → 도착시간 순으로 비교.')
note(s,'예선의 역비례 점수 예시를 일관된 역방향 min–max 점수로 정리했다. 서비스는 수량·기한 제약으로 보장 여부를 먼저 검사하고, 점수에는 결정 시점의 도착 시간 여유만 반영한다. 정시 도착 확률을 뜻하지 않는다. 후보 집합 변화에 따른 상대 점수 변화에 주의한다.')

s=base(q,2,'숫자 검증: 입력부터 선택까지 재계산할 수 있습니다.',
       '설명용 가정: 1,000매 · D=8시간 · B=10,000kgCO₂e · Eₜ=8,200kgCO₂e',appendix=True)
headers=['운송안','비용점수','탄소점수','시간여유','비용 모드','균형 모드','탄소 모드','긴급 모드']
xs=[.85,2.12,3.98,5.83,7.70,9.56,11.41,13.26];widths=[1.02]+[1.68]*7
for x,w,h in zip(xs,widths,headers):tx(s,x,3.01,w,.36,h,16,GRAY,True,PP_ALIGN.CENTER)
for i,a in enumerate('ABC'):
    y=3.75+i*.91;box(s,.75,y,14.51,.71,WHITE if i<2 else PALE)
    vals=[a]+[f'{v:.3f}' for v in DATA['normalized_scores'][i]]+[f"{DATA['RPS'][n][i]:.4f}".rstrip('0').rstrip('.') for n,_ in MODES]
    for x,w,v in zip(xs,widths,vals):tx(s,x,y+.16,w,.37,v,21,TEAL if i==2 else INK,True,PP_ALIGN.CENTER)
box(s,.76,6.89,14.50,.98,NAVY)
tx(s,1.02,7.13,14.00,.48,'C의 실제 배출도 200kg라면  →  Eₜ₊₁ = 8,400kg  →  다음 CUR = 84%',24,WHITE,True)
foot(s,'예상 배출은 선택에 사용하고, 실제 누적 배출은 실행 실적에 따라 갱신합니다. 예상 200kg을 실측값으로 간주하지 않습니다.')
note(s,'전체 계산은 model.py, scenario.csv, model_results.json에 동일하게 제공한다. Emergency C는 0.8625, B는 0.685, A는 0.125다. 비용·배출 입력은 실제 견적이나 배출계수로부터 산출된 것이 아니다.')

s=base(q,3,'가중치를 조금 바꾸면, 선택은 어디에서 달라질까요?',
       '동일 가상 후보에서 γ=0.2를 고정하고, 비용 α와 탄소 β=0.8−α 사이의 민감도를 확인했습니다.',appendix=True)
s.shapes.add_picture(str(OUT/'figures/weight_sensitivity.png'),Inches(.80),Inches(2.97),width=Inches(10.1),height=Inches(4.72))
box(s,11.22,3.02,4.03,4.65,WHITE,LINE)
tx(s,11.48,3.32,3.51,.38,'선택 경계',23,INK,True)
for yy,title,body in [(4.07,'α < 0.2524','C가 우선'),(5.07,'0.2524 < α < 0.6421','B가 우선'),(6.08,'α > 0.6421','A가 우선')]:
    tx(s,11.49,yy,3.50,.32,title,15,GRAY);tx(s,11.49,yy+.40,3.50,.39,body,22,TEAL,True)
foot(s,'경계에서 두 후보가 동점. 이 그림은 가중치가 선택에 미치는 영향을 보여 주며 최적 가중치를 입증하지 않습니다. γ=0.1인 긴급 모드는 제외.')
note(s,'C=B 경계는 α=0.265/1.05, B=A 경계는 α=0.61/0.95. 가상 입력에 의존하는 민감도 분석이다. 더 많은 후보와 현장 비용·배출·납기 자료로 안정성을 확인해야 한다.')

s=base(q,4,'현장 적용에서 지켜야 할 세 가지 경계',
       '탄소예산을 지키려다가 공급을 자동 중단하거나, 모형의 추천을 곧바로 실행하지 않습니다.',appendix=True)
for y,n,a,b in [(3.00,'01','실행 가능한 후보가 없다면','수량·기한·적재·예산을 모두 만족하는 안이 없음을 알리고,\n담당자가 대체 조달·기한 협의·예산 예외 승인 여부를 결정합니다.'),
               (4.60,'02','예상 배출과 실제 배출이 다르면','실적을 갱신하고 다음 선택을 다시 계산합니다. 예산 준수는 보장되지 않습니다.\n실증에서 배출 추정 오차와 예산 초과를 함께 측정합니다.'),
               (6.20,'03','월초 50%와 월말 50%는 다르므로','기본 CUR만으로 남은 기간과 수요를 표현할 수 없습니다.\n기말 예상 배출 점검을 후속 확장으로 두고, 기본 모형과 분리해 검증합니다.')]:
    circle(s,.81,y+.06,.53,TEAL);tx(s,.81,y+.15,.53,.30,n,15,WHITE,True,PP_ALIGN.CENTER)
    tx(s,1.66,y,13.2,.42,a,25,INK,True);tx(s,1.69,y+.65,13.2,.77,b,19,GRAY)
foot(s,'동시 배차에서는 미정산 확정 운행의 예상 배출을 별도로 관리해야 합니다. 이번 계산 예시는 한 번의 순차 의사결정만 다룹니다.')
note(s,'긴급 모드에서 수요 충족률을 낮추거나 장거리 운송을 일괄 중단하던 원안 표현은 제외했다. 자동 중단 대신 불가능한 제약 조합을 명시하고 담당자가 예외를 승인한다. 수요·차량·규격·예산 예외를 모형이 임의 변경하지 않는다.')

s=base(q,5,'검증 설계: 무엇을 비교하고, 무엇을 아직 주장하지 않는가',
       '개념 제안 → 가상 계산 → 과거 요청 재생 → 현장 추천 → 제한적 실행의 순서로 검증합니다.',appendix=True)
box(s,.75,2.99,7.16,4.98,WHITE,LINE);tx(s,1.03,3.27,6.58,.39,'공정한 비교 설계',24,INK,True)
for yy,a,b in [(4.05,'대조군','비용 우선 / 고정 가중치 / 적응형 CBP'),(5.00,'동일 조건','수요·후보 생성 규칙·배출 산정·제약·정보'),(5.95,'공통 기록','비용·실제 배출·정시 공급·기말 예산 초과'),(6.90,'민감도','B·임계값·가중치·수요·배출 오차 변화')]:
    tx(s,1.05,yy,1.34,.34,a,17,TEAL,True);tx(s,2.53,yy,5.07,.63,b,17,GRAY)
box(s,8.30,2.99,6.95,4.98,PALE);tx(s,8.58,3.27,6.39,.39,'현재 제시할 수 있는 근거',24,TEAL,True)
tx(s,8.61,4.05,6.31,1.23,'원안 CBP의 일관된 작동 규칙\n공개 입력으로 재현 가능한 선택 결과\n가중치 변화에 대한 민감도',20,INK)
line(s,8.61,5.64,14.94,5.64,'B8D7D1',1)
tx(s,8.61,5.99,6.29,1.41,'아직 검증되지 않은 것\n현장 감축률 · 비용 절감률 · 전역 최적성\n선행 연구·기존 제품 대비 신규성',20,INK)
foot(s,'자료 근거: 제출한 예선 PPT(아이디어), KPCA 합격 안내 및 배너 예시(제출 형식). 외부 논문·수상작 원문 비교는 이번 자료에서 미검증.')
note(s,'과거 요청 재생 시 각 정책의 선택 이후 재고·누적 배출을 독립적으로 갱신해야 한다. 후보 생성 규칙과 사용 가능한 정보는 동일하게 맞춘다. 수요나 미래 운행의 정답을 한 정책에만 미리 제공하지 않는다. 일정 기간 전체를 반복 평가한 뒤 효과를 말한다.')
save_clean(q,OUT/'cbp-research-v3.pptx')

script='# CBP 본선 4분 발표 대본\n\n목표 시간은 4분입니다. 실제 발표 속도에 맞춰 반드시 소리 내어 리허설하세요. 연구 보충자료 5장은 본 발표에 포함하지 않습니다.\n\n'
for time,title,body in SCRIPTS:script+='## '+time+' | '+title+'\n\n'+body+'\n\n'
(OUT/'발표대본_4분.md').write_text(script,encoding='utf-8')
print('Built: main 6 / exact excerpt 4 / research appendix 5')
