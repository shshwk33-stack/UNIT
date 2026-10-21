"""CBP V4: five-slide argument, editable scientific diagrams, exact banner excerpt."""
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from model import select, State, PLANS, verify
from dataclasses import replace
import math

OUT=Path(__file__).parent
FONT='Noto Sans CJK KR'; DISPLAY='Noto Sans CJK KR'
INK='14243C'; NIGHT='101D31'; PANEL='1C2D46'; WHITE='FFFFFF'; PAPER='F4F6F9'
MINT='54DFC0'; TEAL='087F70'; BLUE='4976DA'; GOLD='E5B66B'; GRAY='687D94'; MUTED='ADBED2'; LINE='DAE3EE'; PALE='E4F4EE'; RED='A34C36'
def new():
    p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9)
    p.core_properties.author='두런두런';p.core_properties.title='CBP | 다음 운송까지 보는 파렛트 탄소예산'
    p.core_properties.subject='제13회 ULS 창의성 공모전 본선 · 원안 기반 4분 발표'
    return p
def box(s,x,y,w,h,c,linec=None,r=0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c)
    if r:sh.adjustments[0]=r
    if linec:sh.line.color.rgb=RGBColor.from_string(linec);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    return sh
def tx(s,x,y,w,h,t,size=20,c=INK,b=False,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,st in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=st;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=b;p.font.color.rgb=RGBColor.from_string(c);p.space_after=Pt(0);p.line_spacing=1.13
        if align is not None:p.alignment=align
    return sh
def ln(s,x1,y1,x2,y2,c=LINE,w=1.2,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(c);sh.line.width=Pt(w)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh
def circle(s,x,y,d,c,linec=None):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c)
    if linec:sh.line.color.rgb=RGBColor.from_string(linec);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    return sh
def arrow(s,x,y,w,c=TEAL,direction='right'):
    kind=MSO_SHAPE.DOWN_ARROW if direction=='down' else MSO_SHAPE.LEFT_ARROW if direction=='left' else MSO_SHAPE.RIGHT_ARROW
    sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(.19 if direction=='down' else w),Inches(w if direction=='down' else .19));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();return sh
def polygon(s,points,c):
    b=s.shapes.build_freeform(Inches(points[0][0]),Inches(points[0][1]));b.add_line_segments([(Inches(x),Inches(y)) for x,y in points[1:]],close=True);sh=b.convert_to_shape();sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();return sh
def tag(s,x,y,w,t,dark=False):box(s,x,y,w,.35,PANEL if dark else PALE);tx(s,x+.10,y+.055,w-.20,.24,t,12,MINT if dark else TEAL,True)
def base(p,n,section,title='',sub='',dark=False,appendix=False):
    s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(NIGHT if dark else PAPER)
    tx(s,.68,.35,3.6,.28,'CBP  /  FORECAST & DECISION',12,MINT if dark else TEAL,True)
    tx(s,6.48,.35,8.85,.28,section.upper(),11,MUTED if dark else GRAY,False,PP_ALIGN.RIGHT)
    if title:tx(s,.68,1.02,14.64,1.19,title,36,WHITE if dark else INK,True)
    if sub:tx(s,.71,2.28,14.55,.54,sub,18,MUTED if dark else GRAY)
    ln(s,.68,8.44,15.32,8.44,'32465F' if dark else LINE,1)
    tx(s,.69,8.61,13.64,.23,'두런두런  ·  제13회 유닛로드시스템(ULS) 창의성 공모전'+('  /  질의응답 보충' if appendix else ''),10,MUTED if dark else GRAY)
    tx(s,14.43,8.55,.88,.35,f'{n:02}',16,MINT if dark else TEAL,True,PP_ALIGN.RIGHT)
    return s
def foot(s,t,dark=False):tx(s,.72,8.06,14.56,.24,t,10,MUTED if dark else GRAY)
def note(s,t):s.notes_slide.notes_text_frame.text=t
def save(p,path):
    for s in p.slides:
        for sh in s.shapes:
            for ref in sh._element.xpath('./p:style/a:effectRef'):ref.set('idx','0')
    p.save(path)
def excerpt(p,indices,path):
    dest=new()
    for i in indices:
        src=p.slides[i];dst=dest.slides.add_slide(dest.slide_layouts[6]);dst.background.fill.solid();dst.background.fill.fore_color.rgb=src.background.fill.fore_color.rgb
        for sh in src.shapes:dst.shapes._spTree.insert_element_before(deepcopy(sh.element),'p:extLst')
        note(dst,src.notes_slide.notes_text_frame.text)
    save(dest,path)
def segmented_ring(s,x,y,d,pct):
    n=50;r=d/2
    for i in range(n):
        a=-math.pi/2+i*2*math.pi/n;xx=x+r+math.cos(a)*r;yy=y+r+math.sin(a)*r
        ln(s,xx,yy,x+r+math.cos(a)*(r-.13),y+r+math.sin(a)*(r-.13),MINT if i<n*pct/100 else '38506B',4)
def pallet_iso(s,x,y,scale=1):
    # Parallel deck boards, geometric sidewalls and support blocks.
    for k in range(3):
        xx=x+k*.68*scale;yy=y+k*.34*scale
        polygon(s,[(xx,yy),(xx+1.98*scale,yy-.98*scale),(xx+2.37*scale,yy-.79*scale),(xx+.39*scale,yy+.20*scale)],'A5EEE0')
        polygon(s,[(xx,yy),(xx+.39*scale,yy+.20*scale),(xx+.39*scale,yy+.38*scale),(xx,yy+.18*scale)],'43A997')
        polygon(s,[(xx+.39*scale,yy+.20*scale),(xx+2.37*scale,yy-.79*scale),(xx+2.37*scale,yy-.61*scale),(xx+.39*scale,yy+.38*scale)],'72CBB9')
    for k in range(3):
        xx=x+.20*scale+k*.68*scale;yy=y+.42*scale+k*.34*scale
        polygon(s,[(xx,yy),(xx+.25*scale,yy+.12*scale),(xx+.25*scale,yy+.47*scale),(xx,yy+.35*scale)],'2E8B7E')
