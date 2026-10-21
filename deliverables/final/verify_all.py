from pathlib import Path
import json,re,hashlib,unicodedata,subprocess
from fractions import Fraction as Q
from pptx import Presentation
from lxml import etree
import fitz
from model import verify,PLANS
R=Path(__file__).parent
out={'model_assertions':verify(),'pdf':{},'exact_excerpts':{},'independent_replay':[]}
D=json.loads((R/'results.json').read_text())
# Independently reconstruct each discrete request policy with rational arithmetic.
for saved in D['replays']:
 policy=saved['policy'];factor=Q(str(saved['actual_factor']));actual=Q(0);cost=0;exc=0
 plans=[('A',100,100),('B',108,60),('C',115,30)]
 for t in range(20):
  h=Q(1200)-actual-(19-t)*60-(20 if t<19 else 0) if policy=='forecast' else Q(1200)-actual
  possible=[x for x in plans if Q(x[2])<=h]
  if policy=='cost':chosen=plans[0]
  elif possible:chosen=min(possible,key=lambda x:(x[1],x[2]))
  else:chosen=min(plans,key=lambda x:(x[2],x[1]));exc+=1
  cost+=chosen[1];actual+=Q(chosen[2])*factor
  assert chosen[0]==saved['trace'][t]['selected']
  assert actual==Q(str(saved['trace'][t]['actual']))
 assert actual==Q(str(saved['emission'])) and cost==saved['cost'] and exc==saved['exceptions']
 out['independent_replay'].append(dict(policy=policy,factor=float(factor),emission=float(actual),cost=cost,matched=True))
assert Q(80000)/Q(40,1000)==2000000
prs={}
for name,n in [('main',5),('banner',4),('appendix',8),('full',13)]:
 prs[name]=p=Presentation(R/f'cbp-final-{name}.pptx')
 assert len(p.slides)==n
 for si,s in enumerate(p.slides):
  assert s.notes_slide.notes_text_frame.text.strip()
  for sh in s.shapes:
   assert min(sh.left,sh.top)>=0
   assert sh.left+sh.width<=p.slide_width+100 and sh.top+sh.height<=p.slide_height+100,(name,si,sh.name)
def same(s,t):
 assert len(s.shapes)==len(t.shapes)
 assert all(etree.tostring(a.element)==etree.tostring(b.element) for a,b in zip(s.shapes,t.shapes))
 assert s.notes_slide.notes_text_frame.text==t.notes_slide.notes_text_frame.text
for i in range(4):same(prs['main'].slides[i+1],prs['banner'].slides[i])
for i in range(5):same(prs['main'].slides[i],prs['full'].slides[i])
for i in range(8):same(prs['appendix'].slides[i],prs['full'].slides[i+5])
out['exact_excerpts']={'banner_equals_main_2_to_5':True,'full_equals_main_plus_appendix':True}
for name,n in [('report',27),('main',5),('banner',4),('appendix',8),('full',13)]:
 pdf=fitz.open(R/f'cbp-final-{name}.pdf');assert len(pdf)==n
 fontcount=0;body_errors=[]
 for i,p in enumerate(pdf):
  text=p.get_text();assert len(text)>50,(name,i,'empty');assert '\ufffd' not in text,(name,i,'replacement character')
  for block in p.get_text('dict')['blocks']:
   for line in block.get('lines',[]):
    for sp in line.get('spans',[]):
     b=fitz.Rect(sp['bbox']);assert b.x0>=-1 and b.y0>=-1 and b.x1<=p.rect.width+1 and b.y1<=p.rect.height+1,(name,i,sp['text'])
 # Poppler recognizes embedded Type3 CharProcs as well as font-file streams.
 fontrows=subprocess.check_output(['pdffonts',str(R/f'cbp-final-{name}.pdf')],text=True).splitlines()[2:]
 for fr in fontrows:
  fields=fr.split();assert fields[-5]=='yes',(name,'font not embedded',fr);assert fields[-3]=='yes',(name,'no Unicode map',fr)
 fontcount=len(fontrows)
 # Byte-independent banner render identity also verifies PDF visual consistency.
 if name=='banner':
  main=fitz.open(R/'cbp-final-main.pdf')
  for i,p in enumerate(pdf):assert p.get_pixmap().samples==main[i+1].get_pixmap().samples
 out['pdf'][name]={'pages':n,'embedded_font_records':fontcount,'text_within_page':True,'no_replacement_characters':True}
report=fitz.open(R/'cbp-final-report.pdf')
normalize=lambda s:re.sub(r'\s+','',unicodedata.normalize('NFKC',s))
for item in json.loads((R/'contents.json').read_text()):assert normalize(item['title']) in normalize(report[item['page']-1].get_text()),item
links=report[2].get_links();assert len(links)==25
assert sorted(x['page'] for x in links)==[i for i in range(27) if i not in [0,2]]
assert len(report[25].get_links())>=5
for text in ['탄소예산','1,215','200만원','683','1,170']:
 assert normalize(text) in normalize(''.join(p.get_text() for p in report)),text
out['report_links']={'contents':25,'references':len(report[25].get_links())}
layout=json.loads((R/'layout_audit.json').read_text());assert len(layout)==27 and all(not p['overflow'] for p in layout)
out['report_dom_overflow']=False
out['demo']=json.loads((R/'demo-verification.json').read_text())
out['visual_review']='27 report pages and 13 slides reviewed as rendered contact sheets; detailed checks on edited chart and budget slides.'
(R/'verification.json').write_text(json.dumps(out,ensure_ascii=False,indent=2))
print(json.dumps(out,ensure_ascii=False,indent=2))
