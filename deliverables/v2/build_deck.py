"""Editable vector deck: four-minute story + exact banner excerpt.
All operational quantities are explicitly illustrative, not measured outcomes.
"""
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

OUT=Path(__file__).parent
W,H=16,9
FONT='Noto Sans CJK KR'
BLACK='172825'; DARK='213B34'; GREEN='16745C'; LIME='DDF485'
CREAM='F5F3EB'; WHITE='FFFFFF'; GRAY='69756E'; LINE='D9DED3'
ORANGE='F36A42'; PALE='E7EEDD'; MUTED='B8C7BD'; WARM='FBE9DF'

def new_prs():
    p=Presentation();p.slide_width=Inches(W);p.slide_height=Inches(H)
    p.core_properties.author='두런두런';p.core_properties.subject='제13회 ULS 창의성 공모전 본선'
    return p

def box(s,x,y,w,h,color=WHITE,r=0,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color)
    if line: sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(1)
    else: sh.line.fill.background()
    if r: sh.adjustments[0]=r
    # No inherited theme shadows: everything remains editable and flat.
    effects=OxmlElement('a:effectLst');sh._element.spPr.append(effects)
    return sh

def tx(s,x,y,w,h,content,size=20,color=BLACK,bold=False,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,t in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_before=Pt(0);p.space_after=Pt(3);p.line_spacing=1.12
        if align is not None:p.alignment=align
    return sh

def line(s,x1,y1,x2,y2,color=LINE,width=2,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh

def circle(s,x,y,d,color):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();return sh

def arrow(s,x,y,w,color=GREEN):
    sh=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(.25));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();return sh

def down_arrow(s,x,y,h,color=ORANGE):
    sh=s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,Inches(x),Inches(y),Inches(.30),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();return sh

def chip(s,x,y,w,content,color=PALE,ink=GREEN,size=12):
    box(s,x,y,w,.36,color,.12);tx(s,x+.10,y+.055,w-.20,.25,content,size,ink,True)

def truck(s,x,y,scale=1,color=GREEN,load=True):
    box(s,x,y,1.25*scale,.61*scale,color,.06)
    box(s,x+1.26*scale,y+.20*scale,.46*scale,.41*scale,color,.06)
    box(s,x+1.33*scale,y+.24*scale,.21*scale,.14*scale,'D7E7E3')
    for xx in [.25,1.30]:circle(s,x+xx*scale,y+.51*scale,.25*scale,BLACK)
    if load:
        for xx in [.12,.45,.78]:box(s,x+xx*scale,y+.14*scale,.25*scale,.28*scale,'C4D9CE',.03)

def pallet(s,x,y,scale=1,count=3):
    for i in range(count):
        yy=y-i*.16*scale
        box(s,x,yy,1.20*scale,.09*scale,'B58750')
        for xx in [.05,.51,1.0]:box(s,x+xx*scale,yy+.09*scale,.13*scale,.07*scale,'8E603A')

def warehouse(s,x,y,scale=1,ink=GREEN):
    box(s,x,y+.23*scale,1.65*scale,1.02*scale,ink,.03)
    roof=s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(x-.08*scale),Inches(y-.25*scale),Inches(1.81*scale),Inches(.65*scale));roof.fill.solid();roof.fill.fore_color.rgb=RGBColor.from_string(ink);roof.line.fill.background()
    box(s,x+.38*scale,y+.66*scale,.85*scale,.59*scale,'C4D9CE')
    for dy in [.77,.90,1.03]:line(s,x+.40*scale,y+dy*scale,x+1.21*scale,y+dy*scale,ink,1)

def header(s,n,kicker,title,subtitle=None,dark=False):
    bg=BLACK if dark else CREAM;s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(bg)
    tx(s,.65,.30,12,.30,'CBP  /  탄소예산 기반 파렛트 회수예약',12,LIME if dark else GREEN,True)
    tx(s,14.30,.30,1.05,.30,f'{n:02} / 06',12,MUTED if dark else GRAY,True,PP_ALIGN.RIGHT)
    tx(s,.65,.88,14.65,1.19,title,34,WHITE if dark else BLACK,True)
    if subtitle:tx(s,.68,2.12,14.55,.48,subtitle,17,MUTED if dark else GRAY)
    line(s,.65,8.48,15.35,8.48,'395348' if dark else LINE,1)
    tx(s,.66,8.63,13.8,.20,'두런두런  ·  제13회 유닛로드시스템(ULS) 창의성 공모전  ·  '+kicker,10,MUTED if dark else GRAY)

def slide(p,n,kicker,title,subtitle=None,dark=False):
    s=p.slides.add_slide(p.slide_layouts[6]);header(s,n,kicker,title,subtitle,dark);return s

SCRIPTS=[
('00:00–00:30','16시 출고, 빈 파렛트 160매',
 '오후 네 시, 출고해야 할 물건은 준비됐는데 빈 파렛트 160매가 부족하다면 어떨까요? 다른 거점에는 파렛트가 있고, 돌아오는 배송차도 있습니다. 하지만 너무 늦게 알면 전용 회수차를 불러야 합니다. 저희 CBP는 필요한 시점에 맞춰 파렛트가 돌아올 자리를 미리 예약하는 아이디어입니다.'),
('00:30–01:15','파렛트·빈칸·탄소를 함께 예약',
 '핵심은 이 회수예약표 한 장입니다. 먼저 다음 출고에 필요한 수량과 도착 시각을 확인합니다. 그다음 회수할 파렛트와, 실제로 실을 수 있는 차량 공간을 함께 예약합니다. 파렛트만 잡아도, 차량만 잡아도 회수는 되지 않기 때문입니다. 여기에 예상 배출을 탄소예산에 예약해 두고, 담당자가 승인합니다. 처음에는 같은 풀링사의 호환 파렛트와 확인된 운행만 대상으로 시작합니다.'),
('01:15–02:00','가상 사례: 100매 + 60매',
 '이해를 위한 가상 사례입니다. 열한 시에 네 시까지 160매가 필요하다는 사실을 확인합니다. 이미 예정된 복귀편 A에 100매, B에 60매를 예약합니다. 열두 시 반에 100매, 한 시 반에 60매가 실제로 도착하면 회수가 끝납니다. 이 조건에서는 전용 회수차를 추가로 보내지 않아도 됩니다. 다만 이 숫자는 실제 성과가 아니라 작동 순서를 설명하는 예시입니다.'),
('02:00–02:40','탄소는 출발 전 선택 조건',
 '탄소예산은 회수 방법을 고를 때 씁니다. 이미 쓴 양과 확정된 예약을 빼고 남은 양을 확인합니다. 화면의 가상 조건에서 전용차는 12만원에 90킬로그램, 복귀편 조합은 13만5천원에 65킬로그램입니다. 남은 예산이 70이면 복귀편 조합이 맞습니다. 물론 비용이 더 듭니다. 운영자가 허용한 비용 상한 안에서 비교하고, 만족하는 안이 없으면 승인 요청을 냅니다. 무조건 비용과 탄소가 함께 줄어든다고 주장하지 않습니다.'),
('02:40–03:25','예약과 실제 도착을 구분',
 '가장 중요한 장치는 이 열네 시입니다. 네 시에 필요하고 대체 회수에 한 시간 반, 여유 시간 삼십 분이 필요하다면 열네 시에는 결정해야 합니다. 이때 예약됐다는 표시가 아니라 실제 들어온 수량을 봅니다. 100매만 들어왔다면 부족한 60매만 대체 회수합니다. 단, 대체 재고와 차량을 사전에 확보할 수 있는 거점에서 적용해야 합니다. 그래야 복귀편을 기다리다가 출고까지 늦어지는 위험을 줄일 수 있습니다.'),
('03:25–04:00','작게 시작하고 세 가지로 검증',
 '첫 적용은 동일 풀링사의 두세 거점이면 됩니다. 출고 계획과 재고, 복귀편 정보를 연결하고 담당자가 추천을 승인하는 방식으로 시작합니다. 검증할 것은 제시간에 필요한 수량이 왔는지, 모든 추가 비용을 포함해 얼마가 들었는지, 실제 운행 배출이 얼마나 달라졌는지입니다. CBP는 빈 파렛트를 급해진 뒤 찾는 대신, 돌아올 시간부터 예약하자는 제안입니다.')
]

p=new_prs()

# 1. A concrete opening scene, not an abstract title page.
s=slide(p,1,'문제와 제안','16시 출고.\n물건은 있는데, 파렛트가 없다면?',dark=True)
chip(s,.68,2.35,2.07,'설명용 가상 장면',DARK,LIME)
tx(s,.72,3.08,5.7,.42,'다음 출고에 필요한 빈 파렛트',20,MUTED)
tx(s,.65,3.61,3.15,1.35,'160',88,LIME,True)
tx(s,3.92,4.17,1,.62,'매',34,WHITE,True)
tx(s,.72,5.31,5.45,1.02,'늦게 알수록\n회수 방법은 줄어듭니다.',26,WHITE,True)
box(s,7.02,2.70,8.31,4.46,DARK,.04)
warehouse(s,7.61,3.38,1.03,'759988');pallet(s,7.82,5.19,.8,4)
tx(s,7.41,5.78,2.4,.35,'다른 거점의 파렛트',17,WHITE,True,PP_ALIGN.CENTER)
truck(s,10.32,3.76,1.19,LIME,False)
tx(s,10.09,4.94,3.2,.77,'배송을 마치고\n돌아오는 차량',18,WHITE,True,PP_ALIGN.CENTER)
arrow(s,9.42,4.20,.67,LIME);arrow(s,12.81,4.20,.70,LIME)
box(s,13.76,3.37,.99,1.65,'567A66',.07)
tx(s,13.77,3.88,.97,.60,'출고\n거점',17,WHITE,True,PP_ALIGN.CENTER)
tx(s,7.47,6.52,7.48,.35,'둘을 미리 연결하면, 회수의 선택지가 생깁니다.',20,LIME,True)
line(s,.72,7.58,15.22,7.58,'395348',1)
tx(s,.72,7.77,14.5,.42,'CBP  |  빈 파렛트가 돌아올 시간과 자리를 예약합니다.',24,LIME,True)

# 2. A tangible reservation ticket with the three things it commits.
s=slide(p,2,'핵심 아이디어','파렛트·빈칸·탄소를\n회수예약표 한 장으로 묶습니다.',
        '재고와 차량을 동시에 확보하고, 예약 시점부터 예상 배출을 반영합니다.')
left=[('01','파렛트 예약','수량 · 규격 · 인계 가능 시각'),('02','빈칸 예약','실을 수 있는 수량 · 도착 시각'),('03','탄소 예약','예상 추가 배출 · 남은 예산')]
for i,(n,a,b) in enumerate(left):
    y=2.94+i*1.34
    circle(s,.72,y+.11,.57,GREEN);tx(s,.72,y+.21,.57,.26,n,14,WHITE,True,PP_ALIGN.CENTER)
    tx(s,1.54,y,4.82,.44,a,25,BLACK,True);tx(s,1.56,y+.59,4.83,.35,b,17,GRAY)
    if i<2:line(s,1.00,y+.75,1.00,y+1.34,LINE,2)
arrow(s,6.62,4.46,.83,GREEN)
box(s,7.85,2.84,7.48,4.73,WHITE,.035,line=LINE)
box(s,7.85,2.84,7.48,.69,BLACK,.035)
tx(s,8.12,3.02,5.65,.32,'CBP  /  회수예약표',20,LIME,True)
chip(s,13.02,3.01,1.98,'풀링사 담당자 승인',LIME,BLACK,11)
ticket=[('필요 물량','160매 · 16:00까지'),('확보할 운행','A 100매 / B 60매'),('실제 도착 목표','12:30 / 13:30'),('대체회수 결정','14:00 · 실입고 확인'),('탄소 처리','확정 예약분을 예산에서 차감')]
for i,(a,b) in enumerate(ticket):
    yy=3.83+i*.61
    tx(s,8.16,yy,1.88,.34,a,16,GRAY);tx(s,10.25,yy,4.66,.37,b,18,BLACK,True)
    if i<4:line(s,8.16,yy+.44,14.98,yy+.44,LINE,1)
chip(s,.73,7.76,14.52,'첫 적용 범위  ·  동일 풀링사 / 호환 규격 / 적재 안전 확인 / 대체회수 가능 거점',PALE,GREEN,16)

# 3. The entire normal case drawn as a timeline.
s=slide(p,3,'회수 시나리오','11시에 예약하면,\n16시 출고를 이렇게 준비합니다.',
        '복귀편은 대체회수 결정 시각인 14시보다 먼저 도착하도록 예약합니다. (설명용 가상 사례)')
events=[(1.15,'11:00','필요량 확인','160매 예약',BLACK),(5.17,'12:30','복귀편 A 도착','100매 확보',GREEN),(9.18,'13:30','복귀편 B 도착','160매 확보',GREEN),(13.21,'16:00','다음 출고','필요량 충족',BLACK)]
line(s,1.45,4.12,14.18,4.12,LINE,7)
for i,(x,t,a,b,c) in enumerate(events):
    tx(s,x-.25,2.92,2.65,.65,t,35,c,True,PP_ALIGN.CENTER)
    circle(s,x+.72,3.96,.32,c)
    tx(s,x-.26,4.54,2.73,.37,a,20,BLACK,True,PP_ALIGN.CENTER)
    tx(s,x-.26,5.04,2.73,.43,b,22,c,True,PP_ALIGN.CENTER)
    if i==1:truck(s,x+.15,5.83,.75,GREEN)
    elif i==2:truck(s,x+.15,5.83,.75,GREEN)
    elif i==0:pallet(s,x+.37,6.39,.75,3)
    else:warehouse(s,x+.18,5.83,.66,BLACK)
chip(s,4.82,3.65,2.17,'+ 100매',PALE,GREEN,15)
chip(s,8.96,3.65,2.17,'+ 60매',PALE,GREEN,15)
line(s,11.91,3.90,11.91,6.71,ORANGE,1.3,True)
tx(s,11.03,6.70,2.10,.25,'14:00 부족분 판단',12,'9B3A1C',True,PP_ALIGN.CENTER)
box(s,.73,7.11,14.55,.93,BLACK,.05)
tx(s,1.00,7.32,11.5,.40,'이 조건에서는 전용 회수차를 추가 배차하지 않아도 됩니다.',23,LIME,True)
tx(s,13.17,7.22,1.65,.54,'0회',32,LIME,True,PP_ALIGN.RIGHT)
tx(s,.75,8.14,14.45,.20,'0회는 이 가상 시나리오의 전용 추가 배차 횟수입니다. 복귀편의 우회·추가 적재·대기 비용과 배출은 별도로 산정합니다.',10,GRAY)

# 4. Carbon stops being an abstract acronym: show the wallet and choices.
s=slide(p,4,'탄소예산의 역할','출발하기 전에,\n남은 탄소예산으로 회수 방법을 고릅니다.',
        '수량·기한·비용 상한을 충족하는 후보 중, 탄소예산 안에서 회수비가 낮은 조합을 선택합니다.')
chip(s,.73,2.80,4.68,'운임·배출량은 모두 설명용 가정',WARM,'9B3A1C',14)
box(s,.73,3.43,4.16,3.96,BLACK,.045)
tx(s,1.04,3.72,3.55,.36,'이번에 쓸 수 있는 양',20,MUTED,True)
tx(s,1.02,4.27,2.59,1.11,'70',72,LIME,True)
tx(s,3.20,4.85,1.34,.41,'kgCO₂e',17,LIME,True)
tx(s,1.05,5.81,3.55,.87,'총예산 300 − 기배출 150\n− 확정 예약 80',18,WHITE)
tx(s,1.05,6.88,3.55,.26,'계산은 기간·단위를 동일하게',12,MUTED)
for x,tag,title,money,emission,sub,bg,accent in [
    (5.31,'비용만 보면','전용 회수차','12만원','90','예산 20kg 초과',WHITE,ORANGE),
    (10.34,'예산을 지키려면','복귀편 A + B','13.5만원','65','예산 5kg 남음',PALE,GREEN)]:
    box(s,x,3.43,4.92,3.96,bg,.045)
    tx(s,x+.26,3.72,4.40,.29,tag,14,accent,True)
    tx(s,x+.26,4.17,4.40,.43,title,25,BLACK,True)
    tx(s,x+.27,4.89,4.40,.58,money,33,BLACK,True)
    tx(s,x+.27,5.59,2.33,.77,emission,48,accent,True)
    tx(s,x+2.28,6.00,2.32,.34,'kgCO₂e',18,accent,True)
    tx(s,x+.27,6.77,4.38,.34,sub,18,accent,True)
tx(s,.78,7.73,14.5,.39,'이 가상 입력의 차이  ·  추가 비용 +1.5만원 / 배출 −25kgCO₂e',22,GREEN,True)
tx(s,.78,8.18,14.5,.18,'기존 배송 본선은 양안에서 동일하게 운행한다고 가정. 비용 상한이 13.5만원 미만이면 예산 준수안은 없으며 담당자 판단이 필요합니다.',10,GRAY)

# 5. Failure is the inventive operational rule: reservations are not arrivals.
s=slide(p,5,'실패 대응','예약된 160매와\n실제로 도착한 160매는 다릅니다.',
        '기다릴 수 있는 마지막 시각을 정하고, 실입고 기준으로 부족분만 대체 회수합니다.')
box(s,.73,2.91,5.06,4.94,BLACK,.045)
tx(s,1.07,3.21,4.35,.41,'대체회수 결정 시각',23,MUTED,True)
tx(s,1.00,3.82,4.42,1.20,'14:00',70,LIME,True)
tx(s,1.08,5.43,4.22,1.32,'필요 시각 16:00\n− 대체 회수 90분\n− 여유 시간 30분',22,WHITE)
tx(s,1.08,7.31,4.20,.26,'운송 소요시간은 현장 입력값',12,MUTED)
box(s,6.20,2.91,9.06,1.12,WHITE,.045)
tx(s,6.50,3.14,2.65,.32,'14시 실입고 확인',18,BLACK,True)
tx(s,9.27,3.09,5.50,.48,'100매 도착 / 60매 미도착',24,ORANGE,True)
down_arrow(s,10.58,4.13,.43,ORANGE)
box(s,6.20,4.70,9.06,1.47,WARM,.045)
tx(s,6.52,4.95,7.95,.48,'부족한 60매만 대체 회수',28,'9B3A1C',True)
tx(s,6.53,5.56,7.95,.31,'예약 상태 확인 → 취소·변경 처리 → 대체 거점·차량 배차',16,BLACK)
tx(s,6.30,6.56,8.73,.82,'대체 재고·차량을 사전 확보할 수 있어야 적용합니다.\n예산 초과가 예상되면 초과량과 함께 승인을 요청합니다.',19,BLACK)
tx(s,6.30,7.65,8.73,.31,'지연된 원래 차량의 물량도 추적해 중복 회수를 정리합니다.',15,GRAY)
tx(s,.77,8.17,14.5,.18,'90분·30분은 가정입니다. 대체 공급이 불가능하거나 소요시간이 늘면 기한 준수를 보장할 수 없으며, 즉시 운영자에게 알립니다.',10,GRAY)

# 6. A strong close with a feasible small starting point and explicit criteria.
s=slide(p,6,'실행과 검증','빈 파렛트도,\n돌아올 시간부터 예약합니다.',dark=True)
tx(s,.72,2.36,14.5,.43,'첫 적용은 동일 풀링사의 2~3개 거점. 출고·재고·복귀편 정보와 담당자 승인으로 시작합니다.',20,MUTED)
cards=[(.73,'01','출고를 지켰나','기한 내 실제 입고량\n납기 미준수 건수'),(5.66,'02','총비용은 얼마인가','운송·우회·대기·상하차\n대체회수·예약 비용 포함'),(10.59,'03','실제 배출은 줄었나','같은 수요·기간·서비스 조건\n전체 운행 경계로 비교')]
for x,n,a,b in cards:
    box(s,x,3.27,4.67,2.79,DARK,.05)
    tx(s,x+.29,3.54,4.09,.37,n,20,LIME,True)
    tx(s,x+.29,4.14,4.09,.48,a,26,WHITE,True)
    tx(s,x+.30,4.99,4.06,.78,b,18,MUTED)
line(s,.77,6.53,15.18,6.53,'395348',1)
tx(s,.77,6.88,14.3,.51,'같이 예약하고 → 실제 도착을 확인하고 → 늦기 전에 부족분만 회수.',27,LIME,True)
tx(s,.77,7.75,14.3,.29,'검증 계획: 현장 규칙 / 같은 정보를 쓴 비용 최소안 / CBP를 구분해 비교합니다.',15,MUTED)
tx(s,.77,8.17,14.3,.18,'현재 단계: 가상 시나리오·계산 검증. 현장 실증 및 선행기술 검토는 추가 과제입니다.',10,MUTED)

for i,s in enumerate(p.slides):
    timing,title,script=SCRIPTS[i]
    s.notes_slide.notes_text_frame.text=f'{timing} | {title}\n\n{script}\n\n주의: 모든 예시 숫자는 합성 가정. 실제 운송·감축 실적이나 독창성 검증 완료로 표현하지 않는다.'
p.core_properties.title='CBP 회수예약표 — 본선 발표 V2'
p.save(OUT/'cbp-main-v2.pptx')

def extract(prs,ids,path,title):
    out=new_prs();out.core_properties.title=title
    for idx in ids:
        source=prs.slides[idx];dest=out.slides.add_slide(out.slide_layouts[6])
        dest.background.fill.solid();dest.background.fill.fore_color.rgb=source.background.fill.fore_color.rgb
        for sh in source.shapes:
            dest.shapes._spTree.insert_element_before(deepcopy(sh.element),'p:extLst')
        dest.notes_slide.notes_text_frame.text=source.notes_slide.notes_text_frame.text
    out.save(path)

extract(p,[1,2,3,4],OUT/'cbp-banner-v2.pptx','CBP 메인 발표 2~5장 발췌 — 배너용 4쪽')

# Separate Q&A deck so the 4-minute main stays focused.
q=new_prs()
s=slide(q,1,'질의응답 참고','백홀 활용 자체보다,\n어떻게 확정하고 언제 포기할지가 핵심입니다.',
        '선행기술과 실제 운영 기능은 추가 확인이 필요합니다. 아래는 제안한 설계의 초점입니다.')
for i,(a,b,c) in enumerate([
    ('동시 예약','재고만 또는 빈칸만 잡히는 상황 방지','파렛트 수량과 차량 용량을 같은 예약번호로 확정'),
    ('실입고 기준 마감','예약 성공과 실제 확보를 구분','대체회수 소요시간을 거꾸로 계산해 재판단'),
    ('탄소 예약 장부','확정 운행의 예상 배출도 미리 반영','취소·변경·실행에 따라 예약량과 기배출량 갱신')]):
    y=3.0+i*1.39;box(s,.75,y,14.48,1.16,WHITE,.05)
    tx(s,1.02,y+.23,2.39,.37,a,23,GREEN,True);tx(s,3.83,y+.14,10.95,.34,b,21,BLACK,True);tx(s,3.84,y+.65,10.95,.31,c,17,GRAY)
tx(s,.79,7.59,14.45,.55,'“세계 최초” 대신, 목표 현장에서 이 세 규칙이 실제로 필요한지 확인하겠습니다.',23,GREEN,True)

s=slide(q,2,'질의응답 참고','감축량은 같은 조건에서,\n비용은 빠짐없이 비교해야 합니다.',
        '사전 정보의 효과와 탄소예산의 효과를 구분하고, 가상 예시를 성과로 바꾸어 말하지 않습니다.')
for x,n,a,b in [( .75,'A','현장 운영 규칙','실제 절차를 조사한 뒤 재현'),(5.65,'B','동일 정보 · 비용 최소','복귀편 정보까지 같은 조건 제공'),(10.55,'C','CBP','B에 탄소예산과 예약 규칙 적용')]:
    box(s,x,3.10,4.66,1.87,WHITE,.05);tx(s,x+.26,3.33,.60,.42,n,26,GREEN,True);tx(s,x+.26,3.91,4.12,.37,a,22,BLACK,True);tx(s,x+.26,4.49,4.12,.25,b,14,GRAY)
tx(s,.81,5.43,14.3,.57,'A → B : 사전 정보·매칭의 효과       B → C : 탄소예산·예약 정책의 효과',23,GREEN,True)
box(s,.75,6.35,14.46,1.48,PALE,.05)
tx(s,1.01,6.57,13.94,.40,'복귀차량도 추가 적재·우회·대기 배출이 생깁니다.',25,GREEN,True)
tx(s,1.02,7.16,13.90,.37,'기존 본선이 같은 경우만 공통분을 상쇄합니다. 회수·대체차·지연물량 정리까지 비용과 배출에 포함합니다.',17,BLACK)

s=slide(q,3,'질의응답 참고','비용이 늘어도 도입할 이유가 있나요?',
        '모든 운영자에게 유리하다고 단정하지 않습니다. 비용 상한과 실제 탄소 요구가 있는지 확인합니다.')
box(s,.75,2.94,6.63,3.24,BLACK,.05);tx(s,1.08,3.30,5.95,.42,'가상 예시의 추가 비용',23,MUTED,True)
tx(s,1.08,4.06,5.93,.74,'1.5만원 ÷ 25kg',37,LIME,True);tx(s,1.09,5.12,5.91,.51,'= 60만원 / tCO₂e',30,WHITE,True)
tx(s,7.91,3.18,7.03,1.04,'이 비율은 시장 가격도,\n현장 감축 비용도 아닙니다.',27,BLACK,True)
tx(s,7.94,4.71,6.99,1.29,'실제 운임·거리·차량 조건으로 다시 계산합니다.\n탄소 요구가 약하고 비용만 늘면\n도입 타당성이 낮아질 수 있습니다.',20,GRAY)
box(s,.75,6.70,14.48,1.10,PALE,.05)
tx(s,1.03,6.94,13.90,.65,'필요한 첫 증거: 실제 회수 1건의 수요·시각·운임·복귀편·연결 실패 이유',23,GREEN,True)

for i,s in enumerate(q.slides):
    # These are reference pages, not numbered as part of the main six.
    for sh in s.shapes:
        if sh.has_text_frame and sh.text==f'{i+1:02} / 06':sh.text=f'Q&A {i+1:02}'
q.core_properties.title='CBP V2 — 질의응답 참고자료'
q.save(OUT/'cbp-qa-v2.pptx')

doc=['# CBP V2 — 4분 발표 대본','',
     '메인 발표는 6장입니다. 배너용은 메인 2~5장을 그대로 발췌했습니다. 아래 시간은 리허설 목표이며 실제 낭독으로 조정해야 합니다. 3분40초 내외 낭독을 목표로 하고 전환·호흡 여유를 남기세요.','']
for i,(time,title,script) in enumerate(SCRIPTS,1):doc.extend([f'## {i}장 · {time} · {title}','',script,''])
doc.extend(['## 발표할 때 지킬 표현','','- “가상 사례입니다”를 3장과 4장에서 말로도 밝힙니다.','- 160매·도착 시각·90분·30분·운임·배출량은 설명용 입력입니다. 실제 업계 평균이나 현장 성과가 아닙니다.','- 배너에 인쇄될 수 있으므로 가상 예시 표시를 삭제하지 않습니다.','- 전용 추가 배차 0회는 제시한 정상 사례 한 건의 결과입니다. 일반적인 절감률이 아닙니다.','- 재고·차량·소요시간이 확보되지 않으면 납기를 보장할 수 없습니다.','- “세계 최초”, “업계에는 없는 기능”은 선행기술·현장 확인 전에는 사용하지 않습니다.',''])
(OUT/'발표대본_V2.md').write_text('\n'.join(doc))
print('Created main (6), banner (4), Q&A (3), and speaker script.')
