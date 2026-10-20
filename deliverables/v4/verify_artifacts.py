"""Check exact excerpts and meaningful numerical behavior."""
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pypdf import PdfReader
from cbp_model import verify, select
import json
root=Path(__file__).parent
def same(a,b):
    assert len(a.shapes)==len(b.shapes)
    assert all(etree.tostring(x.element)==etree.tostring(y.element) for x,y in zip(a.shapes,b.shapes))
    assert a.notes_slide.notes_text_frame.text==b.notes_slide.notes_text_frame.text
docs={}
for name,count in [('main',5),('banner',4),('research',6),('full',11)]:
    p=Presentation(root/f'cbp-{name}-v4.pptx');pdf=PdfReader(root/f'cbp-{name}-v4.pdf');docs[name]=p
    assert len(p.slides)==len(pdf.pages)==count
    for i,s in enumerate(p.slides):
        assert s.notes_slide.notes_text_frame.text.strip()
        for sh in s.shapes:
            assert sh.left>=0 and sh.top>=0
            assert sh.left+sh.width<=p.slide_width+100 and sh.top+sh.height<=p.slide_height+100,(name,i,sh.name)
            for effect in sh._element.xpath('./p:style/a:effectRef'):assert effect.get('idx')=='0'
    print(name,count,'pages; shape bounds and notes passed')
for a,b in zip(list(docs['main'].slides)[1:5],docs['banner'].slides):same(a,b)
for a,b in zip(docs['main'].slides,list(docs['full'].slides)[:5]):same(a,b)
for a,b in zip(docs['research'].slides,list(docs['full'].slides)[5:]):same(a,b)
r=json.loads((root/'model_results.json').read_text());assert r['base']['selected']==select()['selected']
print('Exact excerpts and combined deck passed; numerical checks:',verify())
