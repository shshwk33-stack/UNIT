"""CBP V4: five-slide argument, editable scientific diagrams, exact banner excerpt."""
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from cbp_model import select, State, PLANS, MODES, verify
from dataclasses import replace
import math

OUT=Path(__file__).parent
FONT='Noto Sans CJK KR'; DISPLAY='Noto Sans CJK KR'
INK='14243C'; NIGHT='101D31'; PANEL='1C2D46'; WHITE='FFFFFF'; PAPER='F4F6F9'
MINT='54DFC0'; TEAL='087F70'; BLUE='4976DA'; GOLD='E5B66B'; GRAY='687D94'; MUTED='ADBED2'; LINE='DAE3EE'; PALE='E4F4EE'; RED='A34C36'
def new():
    p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9)
    p.core_properties.author='두런두런';p.core_properties.title='CBP | 탄소예산이 바꾸는 파렛트 재배치'
    p.core_properties.subject='제13회 ULS 창의성 공모전 본선 · 원안 기반 4분 발표'
    return p
def box(s,x,y,w,h,c,linec=None,r=0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c)
    if r:sh.adjustments[0]=r
    if linec:sh.line.color.rgb=RGBColor.from_string(linec);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    return sh
def tx(s,x,y,w,h,t,size=20,c=INK,b=False,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,st in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=st;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=b;p.font.color.rgb=RGBColor.from_string(c);p.space_after=Pt(0);p.line_spacing=1.13
        if align is not None:p.alignment=align
    return sh
def ln(s,x1,y1,x2,y2,c=LINE,w=1.2,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(c);sh.line.width=Pt(w)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh
def circle(s,x,y,d,c,linec=None):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c)
    if linec:sh.line.color.rgb=RGBColor.from_string(linec);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    return sh
def arrow(s,x,y,w,c=TEAL,direction='right'):
    kind=MSO_SHAPE.DOWN_ARROW if direction=='down' else MSO_SHAPE.LEFT_ARROW if direction=='left' else MSO_SHAPE.RIGHT_ARROW
    sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(.19 if direction=='down' else w),Inches(w if direction=='down' else .19));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();return sh
def polygon(s,points,c):
    b=s.shapes.build_freeform(Inches(points[0][0]),Inches(points[0][1]));b.add_line_segments([(Inches(x),Inches(y)) for x,y in points[1:]],close=True);sh=b.convert_to_shape();sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();return sh
def tag(s,x,y,w,t,dark=False):box(s,x,y,w,.35,PANEL if dark else PALE);tx(s,x+.10,y+.055,w-.20,.24,t,12,MINT if dark else TEAL,True)
def base(p,n,section,title='',sub='',dark=False,appendix=False):
    s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(NIGHT if dark else PAPER)
    tx(s,.68,.35,3.6,.28,'CBP  /  DECISION DESIGN',12,MINT if dark else TEAL,True)
    tx(s,6.48,.35,8.85,.28,section.upper(),11,MUTED if dark else GRAY,False,PP_ALIGN.RIGHT)
    if title:tx(s,.68,1.02,14.64,1.19,title,36,WHITE if dark else INK,True)
    if sub:tx(s,.71,2.28,14.55,.54,sub,18,MUTED if dark else GRAY)
    ln(s,.68,8.44,15.32,8.44,'32465F' if dark else LINE,1)
    tx(s,.69,8.61,13.64,.23,'두런두런  ·  제13회 유닛로드시스템(ULS) 창의성 공모전'+('  /  질의응답 보충' if appendix else ''),10,MUTED if dark else GRAY)
    tx(s,14.43,8.55,.88,.35,f'{n:02}',16,MINT if dark else TEAL,True,PP_ALIGN.RIGHT)
    return s
def foot(s,t,dark=False):tx(s,.72,8.06,14.56,.24,t,10,MUTED if dark else GRAY)
def note(s,t):s.notes_slide.notes_text_frame.text=t
def save(p,path):
    for s in p.slides:
        for sh in s.shapes:
            for ref in sh._element.xpath('./p:style/a:effectRef'):ref.set('idx','0')
    p.save(path)
def excerpt(p,indices,path):
    dest=new()
    for i in indices:
        src=p.slides[i];dst=dest.slides.add_slide(dest.slide_layouts[6]);dst.background.fill.solid();dst.background.fill.fore_color.rgb=src.background.fill.fore_color.rgb
        for sh in src.shapes:dst.shapes._spTree.insert_element_before(deepcopy(sh.element),'p:extLst')
        note(dst,src.notes_slide.notes_text_frame.text)
    save(dest,path)
def segmented_ring(s,x,y,d,pct):
    n=50;r=d/2
    for i in range(n):
        a=-math.pi/2+i*2*math.pi/n;xx=x+r+math.cos(a)*r;yy=y+r+math.sin(a)*r
        ln(s,xx,yy,x+r+math.cos(a)*(r-.13),y+r+math.sin(a)*(r-.13),MINT if i<n*pct/100 else '38506B',4)
def pallet_iso(s,x,y,scale=1):
    # Parallel deck boards, geometric sidewalls and support blocks.
    for k in range(3):
        xx=x+k*.68*scale;yy=y+k*.34*scale
        polygon(s,[(xx,yy),(xx+1.98*scale,yy-.98*scale),(xx+2.37*scale,yy-.79*scale),(xx+.39*scale,yy+.20*scale)],'A5EEE0')
        polygon(s,[(xx,yy),(xx+.39*scale,yy+.20*scale),(xx+.39*scale,yy+.38*scale),(xx,yy+.18*scale)],'43A997')
        polygon(s,[(xx+.39*scale,yy+.20*scale),(xx+2.37*scale,yy-.79*scale),(xx+2.37*scale,yy-.61*scale),(xx+.39*scale,yy+.38*scale)],'72CBB9')
    for k in range(3):
        xx=x+.20*scale+k*.68*scale;yy=y+.42*scale+k*.34*scale
        polygon(s,[(xx,yy),(xx+.25*scale,yy+.12*scale),(xx+.25*scale,yy+.47*scale),(xx,yy+.35*scale)],'2E8B7E')

SCRIPTS=[
('00:00–00:30','문제와 제안','안녕하십니까. 두런두런의 CBP, 탄소예산이 바꾸는 파렛트 재배치를 제안합니다. 같은 파렛트 천 매를 보내더라도 비용이 낮은 안과 배출이 적은 안은 다를 수 있습니다. 그렇다면 이번 기간의 탄소예산을 얼마나 썼는지도 다음 선택에 반영할 수 있지 않을까요? CBP는 이 상태에 맞춰 재배치 기준을 바꿉니다.'),
('00:30–01:15','의사결정 순환','먼저 기간 예산 대비 누적 배출로 사용률을 구합니다. 만 킬로그램 중 팔천이백을 썼다면 팔십이 퍼센트입니다. 이 사용률에 따라 운영 모드를 바꾸고, 비용과 탄소에 주는 비중을 조정합니다. 여기서 중요한 순서가 있습니다. 수량과 도착 기한, 호환 규격과 적재, 비용과 탄소 한도를 먼저 검사합니다. 통과한 안끼리 점수를 비교하고, 실행 뒤 실제 배출을 반영합니다. 조건을 만족하는 안이 없으면 담당자 판단을 요청합니다.'),
('01:15–02:10','하나의 계산 사례','가상 사례입니다. 모두 천 매를 여덟 시간 안에 보낼 수 있습니다. A는 백만 원에 배출 사백 킬로그램, B는 백십만 원에 이백육십, C는 백사십만 원에 이백입니다. 사용률 팔십이 퍼센트의 탄소 모드에서는 비용 이십, 탄소 육십, 시간 여유 이십의 비중을 적용합니다. 계산하면 C가 가장 높습니다. 하지만 C는 A보다 사십만 원 비쌉니다. 그래서 비용 한도도 함께 둡니다. 한도가 백이십만 원이라면 C를 제외하고 B를 추천합니다. 이 수치는 성과가 아니라 선택 원리를 설명하는 가정입니다.'),
('02:10–03:00','고정 기준과 상태 연동','같은 세 후보를 유지하고 예산 상태만 바꿔 보겠습니다. 비용 중심 기준은 A를, 고정 가중치는 B를 계속 선택합니다. 반면 CBP는 사용률에 따라 A에서 B, C로 선택을 바꿉니다. 제안의 초점은 비용과 탄소를 함께 계산한다는 사실보다, 기간 예산 상태를 다음 이동의 기준에 연결한다는 데 있습니다. 이 표는 비교 기준을 명시한 개념 실험입니다. 특정 기업의 운영 방식을 단정하거나 기존 기술보다 우수하다고 증명한 결과는 아닙니다.'),
('03:00–03:50','구현과 실증','추천 계산기는 만들었습니다. 사용률과 비용 한도를 바꾸면 추천이 달라지고, 지연되거나 조건을 넘는 안은 제외됩니다. 현장 적용은 동일 풀링사의 호환 파렛트부터 시작하겠습니다. 같은 요청과 정보로 비용 중심 방식, 고정 가중치 방식, CBP를 비교하고, 총비용과 실제 배출, 정시 공급, 기간말 예산 초과를 함께 보겠습니다. 가중치와 경계값, 남은 기간과 수요는 실제 자료로 조정해야 합니다. 탄소를 보고하는 숫자에서, 다음 파렛트 이동을 바꾸는 신호로. 감사합니다.')]

p=new()
# 01 / editorial hero: one promise, one physical object, one state variable.
s=base(p,1,'01 / PROBLEM & PROPOSAL',dark=True)
tx(s,.72,1.23,8.33,2.15,'탄소예산이 바꾸는\n파렛트 재배치',47,WHITE,True)
tx(s,.77,3.81,6.78,.52,'Carbon Budget Pooling',24,MINT,True)
tx(s,.78,4.64,6.65,1.22,'같은 공급 요청에도,\n예산 상태에 맞춰 선택 기준을 바꿉니다.',23,MUTED)
ln(s,.79,6.50,6.88,6.50,'38506B',1)
tx(s,.80,6.84,6.44,.34,'하나의 운영 질문',13,MINT,True)
tx(s,.80,7.30,6.62,.53,'“지금도 가장 싼 안을 고르는 게 맞을까?”',23,WHITE,True)
for xx,yy in [(9.02,3.8),(14.60,3.8),(9.33,6.75),(14.70,6.75)]:
    circle(s,xx,yy,.18,'69879F');ln(s,xx+.09,yy+.09,12.16,5.63,'344D67',1.6)
pallet_iso(s,9.93,5.76,1.10)
segmented_ring(s,10.83,1.19,2.32,82)
tx(s,10.87,1.92,2.23,.64,'82%',43,MINT,True,PP_ALIGN.CENTER)
tx(s,10.87,2.67,2.23,.26,'탄소예산 사용률',12,MUTED,False,PP_ALIGN.CENTER)
tag(s,10.77,7.20,3.80,'예산 상태 → 재배치 우선순위',True)
foot(s,'예선 원안 계승: CUR · 4단계 운영 모드 · 재배치 우선순위 점수(RPS). 그림의 82%는 설명용 상태입니다.',True)

# 02 / the control loop, with a real feasibility gate before the scoring step.
s=base(p,2,'02 / HOW IT WORKS','예산 상태를 읽고, 다음 운송의 기준을 바꿉니다.',
       '탄소예산 사용률(CUR) = 누적 실제 배출 ÷ 기간별 내부 탄소예산 × 100')
box(s,.73,3.03,3.18,4.44,NIGHT)
tx(s,1.00,3.33,2.64,.32,'01  상태 확인',17,MUTED,True)
tx(s,.97,4.00,2.70,1.02,'82%',68,MINT,True)
tx(s,1.02,5.32,2.63,.75,'8,200 / 10,000\nkgCO₂e',19,WHITE)
tx(s,1.02,6.71,2.64,.32,'다음 선택의 입력',16,MUTED)
arrow(s,4.12,4.42,.46)
tx(s,4.93,3.37,4.22,.32,'02  운영 모드',17,TEAL,True)
tx(s,4.90,4.00,4.24,.56,'탄소에 더 큰 비중',27,INK,True)
parts=[(.2,BLUE,'20'),(.6,TEAL,'60'),(.2,GOLD,'20')];x=4.94
for v,c,t in parts:box(s,x,5.02,4.08*v,.50,c);tx(s,x,5.13,4.08*v,.28,t,14,WHITE if c!=GOLD else INK,True,PP_ALIGN.CENTER);x+=4.08*v
tx(s,4.94,5.76,4.08,.33,'비용 20 · 탄소 60 · 시간 20',16,GRAY)
arrow(s,9.46,4.42,.46)
tx(s,10.28,3.37,4.80,.32,'03  가능 후보 → 점수 비교',17,TEAL,True)
for yy,t in [(4.05,'공급 조건  수량·기한·규격·적재'),(4.70,'운영 한도  비용·예상 탄소')]:
    circle(s,10.30,yy+.05,.18,TEAL);tx(s,10.66,yy,4.44,.40,t,19,INK,True)
tx(s,10.30,5.52,4.75,.74,'통과한 안만 종합점수로 비교\n불가능하면 담당자 판단',19,GRAY)
box(s,4.91,6.65,10.36,.82,NIGHT);tx(s,5.17,6.89,9.83,.35,'04  실행 → 실제 배출 갱신 → 다음 요청에 다시 적용',22,WHITE,True)
ln(s,10.91,6.23,10.91,6.61,TEAL,1.5);arrow(s,10.82,6.36,.24,TEAL,'down')
ln(s,4.91,7.08,4.29,7.08,TEAL,1.5);ln(s,4.29,7.08,4.29,6.26,TEAL,1.5);arrow(s,3.91,6.16,.38,TEAL,'left')
foot(s,'사용률은 실제 누적 배출 기준. 탄소 한도는 총예산−누적 실제 배출−확정·미정산 예상 배출로 계산. 내부 한도이며 법정 배출권과 구분합니다.')

# 03 / one concrete example, one recommendation, one cost safeguard.
s=base(p,3,'03 / WORKED EXAMPLE','82%를 썼다면, 어떤 운송안을 고를까요?',
       '동일 요청: 파렛트 1,000매 · 8시간 이내. 비용 한도 140만원. 모든 숫자는 설명용 가정입니다.',True)
tx(s,.83,3.12,6.44,.32,'같은 공급 조건의 세 운송안',19,MUTED,True)
cols=[(.94,1.05,'운송안'),(2.25,1.36,'비용\n만원'),(4.04,1.67,'배출\nkgCO₂e'),(6.04,1.14,'도착\n시간')]
for x,w,t in cols:tx(s,x,3.73,w,.62,t,14,MUTED,True,PP_ALIGN.CENTER)
for i,pl in enumerate(PLANS):
    y=4.60+i*.86
    if i==2:box(s,.80,y-.05,6.58,.72,PANEL)
    for (x,w,_),v in zip(cols,[pl.id,str(pl.cost),str(pl.emission),str(pl.arrival)]):tx(s,x,y+.06,w,.43,v,24,MINT if i==2 else WHITE,True,PP_ALIGN.CENTER)
ln(s,7.76,3.16,7.76,7.62,'3C516C',1)
tx(s,8.28,3.17,6.73,.35,'탄소 모드  /  비용 20 · 탄소 60 · 시간 20',18,MUTED,True)
tx(s,8.25,3.84,1.80,1.47,'C',94,MINT,True)
tx(s,10.37,4.17,4.65,.52,'종합점수 0.725',29,WHITE,True)
tx(s,10.39,4.91,4.65,.34,'A 0.250  /  B 0.670',18,MUTED)
ln(s,8.29,5.78,15.12,5.78,'3C516C',1)
tx(s,8.30,6.04,6.71,.41,'A → C : 비용 +40만원 / 배출 −200kg',22,WHITE,True)
box(s,8.28,6.86,6.87,.70,'25493F')
tx(s,8.51,7.07,6.41,.33,'비용 한도가 120만원이면  →  C 제외, B 추천',18,MINT,True)
foot(s,'RPS는 허용 후보의 비용·배출·시간 여유를 0–1 점수로 맞춘 가중합입니다. 비용 한도 변경 시 정규화 재계산. 현장 감축 성과가 아닙니다.',True)

# 04 / explicit controlled comparison, not a blanket claim about incumbents.
s=base(p,4,'04 / WHAT CHANGES','같은 후보. 다른 예산 상태. 달라지는 선택.',
       '비교 기준을 고정하고, 예산 사용률만 바꾼 단일 의사결정 실험입니다.')
columns=[5.89,8.51,11.13,13.75]
tx(s,.86,3.17,4.40,.37,'비교 정책',18,GRAY,True)
for x,cur in zip(columns,[50,68,82,93]):tx(s,x-.63,3.02,1.80,.55,f'{cur}%',31,TEAL,True,PP_ALIGN.CENTER)
tx(s,12.85,3.68,2.42,.27,'탄소예산 사용률',12,GRAY,False,PP_ALIGN.RIGHT)
rows=[(4.27,'비용 우선','비용만으로 순위 결정','AAAA'),(5.36,'고정 가중치','비용 50 · 탄소 30 · 시간 20','BBBB'),(6.56,'CBP','예산 상태에 따라 가중치 조정','ABCC')]
for yy,title,desc,seq in rows:
    if title=='CBP':box(s,.76,yy-.12,14.50,1.01,PALE)
    tx(s,1.02,yy,3.95,.40,title,25,TEAL if title=='CBP' else INK,True);tx(s,1.03,yy+.50,4.18,.28,desc,13,GRAY)
    for x,a in zip(columns,seq):
        circle(s,x-.28,yy+.02,.58,TEAL if title=='CBP' else WHITE)
        tx(s,x-.28,yy+.11,.58,.36,a,22,WHITE if title=='CBP' else GRAY,True,PP_ALIGN.CENTER)
    if title!='CBP':ln(s,1.02,yy+.86,15.02,yy+.86,LINE,1)
for x in columns[:-1]:arrow(s,x+.56,6.75,.77,TEAL)
foot(s,'세 후보 모두 허용되며 동일한 필수 제약을 적용. 고정 정책은 여기서 정의한 대조군이며 특정 기업·선행 연구의 실제 운영을 대표하지 않습니다.')

# 05 / proof of implementation, credible adoption, memorable close.
s=base(p,5,'05 / PROTOTYPE → PILOT','계산은 구현했습니다.\n다음 검증은 현장 데이터입니다.',dark=True)
tx(s,.80,2.91,5.35,.32,'추천 계산기에서 확인한 동작',16,MINT,True)
for yy,left,right in [(3.65,'사용률 50% → 82%','A → C'),(4.43,'비용 한도 140 → 120만원','C → B'),(5.21,'가능한 후보가 하나도 없으면','담당자 판단')]:
    tx(s,.81,yy,4.56,.37,left,18,MUTED);tx(s,5.25,yy-.03,2.22,.43,right,23,WHITE,True,PP_ALIGN.RIGHT)
    ln(s,.81,yy+.58,7.47,yy+.58,'334B65',1)
ln(s,8.02,2.99,8.02,6.45,'334B65',1)
tx(s,8.58,2.91,6.45,.32,'현장 적용은 작은 범위부터',16,MINT,True)
for i,(yy,a,b) in enumerate([(3.62,'동일 풀링사 · 호환 파렛트','재고·운행·비용·배출 자료 확보'),(4.65,'동일 조건으로 세 정책 비교','비용 우선 / 고정 가중치 / CBP'),(5.68,'네 지표를 함께 확인','총비용 · 실제 배출 · 정시 공급 · 예산 초과')]):
    tx(s,8.60,yy,.43,.32,str(i+1).zfill(2),15,MINT,True);tx(s,9.21,yy-.02,5.93,.40,a,22,WHITE,True);tx(s,9.22,yy+.48,5.91,.34,b,15,MUTED)
box(s,.80,7.02,.08,.68,MINT)
tx(s,1.10,7.03,14.13,.61,'탄소를 보고하는 숫자에서, 다음 파렛트 이동을 바꾸는 신호로.',27,WHITE,True)
foot(s,'개념시연용 프로토타입이며 실제 배차 연동·현장 효과는 미검증. 경계값·가중치와 남은 기간·수요 반영은 실증에서 조정합니다.',True)
for s,(_,title,script) in zip(p.slides,SCRIPTS):note(s,title+'\n\n'+script)
save(p,OUT/'cbp-main-v4.pptx')
excerpt(p,[1,2,3,4],OUT/'cbp-banner-v4.pptx')

# APPENDIX 1 / mathematical contract.
s=base(p,1,'APPENDIX / MODEL','“가능한가?”와 “무엇이 더 좋은가?”를 분리합니다.',
       '하나의 후보 j는 필요한 파렛트를 공급하는 완결 운송안입니다. 합배송 최적 조합 생성은 현재 범위 밖입니다.',appendix=True)
box(s,.78,3.03,6.91,4.77,WHITE)
tx(s,1.04,3.28,6.35,.42,'1. 실행 가능 후보 집합 F',24,TEAL,True)
tx(s,1.06,4.04,6.33,1.48,'공급량·적재 가능량 ≥ 요청량\n도착시간 ≤ 허용시간\n호환 규격·적재 안전 확인',20,INK)
tx(s,1.06,5.81,6.33,1.39,'비용 ≤ 운영자가 정한 비용 한도\n예상 배출 ≤ B − Eₜ − Kₜ\nKₜ: 확정·미정산 운행의 예상 배출',19,INK)
box(s,8.04,3.03,7.20,4.77,PALE)
tx(s,8.31,3.28,6.67,.42,'2. 후보 집합 안에서 평가',24,TEAL,True)
tx(s,8.32,4.03,6.62,1.69,'비용점수 = (Cmax−Cⱼ)/(Cmax−Cmin)\n탄소점수 = (emax−eⱼ)/(emax−emin)\n시간여유 = (D−Tⱼ)/D',19,INK)
tx(s,8.32,6.00,6.61,1.16,'RPSⱼ = α·비용 + β·탄소 + γ·시간여유\nj* = arg maxⱼ∈F RPSⱼ\nα+β+γ=1,  각 가중치 ≥ 0',19,INK)
foot(s,'min·max는 F 안에서 계산. 값이 모두 같으면 해당 항목은 모두 1점. 동점은 배출→비용→시간→ID 순. F가 비면 자동 실행 없이 예외 판단.')
note(s,'Python cbp_model.py에 공급, 적재량, 호환, 적재안전, 기한, 비용, 잔여예산 조건을 구현했다. 비용·배출·시간의 단위는 각각 만원·kgCO2e·시간. 현재는 주어진 후보의 추천 계산이며 실제 재고 검증이나 배차 실행이 아니다.')

# APPENDIX 2 / every score and every original weight.
s=base(p,2,'APPENDIX / REPRODUCIBLE CALCULATION','원안의 네 모드를 유지하고, 점수 계산을 공개했습니다.',
       '상태 구간과 가중치는 설계 가정입니다. 데이터로 검증된 최적값이 아닙니다.',appendix=True)
xs=[.92,4.33,7.79,11.47]
for i,(x,(name,w),rng) in enumerate(zip(xs,MODES,['<60%','60–<75%','75–<90%','≥90%'])):
    tx(s,x,3.10,3.10,.37,name,23,TEAL if i==2 else INK,True)
    tx(s,x,3.65,3.10,.30,'CUR '+rng,16,GRAY)
    pos=x
    for v,c in zip(w,[BLUE,TEAL,GOLD]):box(s,pos,4.23,3.06*v,.46,c);tx(s,pos,4.33,3.06*v,.25,str(round(v*100)),12,WHITE if c!=GOLD else INK,True,PP_ALIGN.CENTER);pos+=3.06*v
    state=replace(State(),used=[5000,6800,8200,9300][i]);r=select(state=state)
    for k,a in enumerate('ABC'):
        yy=5.03+k*.69;tx(s,x,yy,.48,.33,a,19,INK,True);tx(s,x+.72,yy,2.28,.33,f"{r['scores'][a]['RPS']:.4f}".rstrip('0').rstrip('.'),19,TEAL if a==r['selected'] else GRAY,a==r['selected'])
    tag(s,x,7.22,3.06,'선택 '+r['selected'])
foot(s,'정규화 점수: A=(1,0,0.25), B=(0.75,0.70,0.50), C=(0,1,0.625). 순서는 비용·탄소·시간 여유. 비용 한도 140만원으로 세 후보 모두 허용.')
note(s,'C의 Carbon 점수=0.2×0+0.6×1+0.2×0.625=0.725. Cost에서A=.75, Balanced에서B=.685, Emergency에서C=.8625. 모든 숫자는 합성 입력이며 model_results.json과 scenario.csv에서 재현할 수 있다.')

# APPENDIX 3 / the added operational safeguard is tested.
s=base(p,3,'APPENDIX / EXCEPTIONS','탄소 비중을 높여도, 비용·공급 조건은 넘기지 않습니다.',
       '정책이 바꾸는 것은 후보의 순위입니다. 운영 한도나 필수 공급 조건을 임의로 낮추지 않습니다.',appendix=True)
for x,n,title,sub,result in [(.85,'01','비용 한도 120만원','CUR 82%, 세 후보 중 C는 140만원','C 제외 → B 추천'),(5.81,'02','C 도착이 9시간','허용 8시간, C만 지연 발생','C 제외 → B 추천'),(10.77,'03','잔여예산 100kg','최소 배출 후보도 200kg 필요','후보 없음 → 담당자 판단')]:
    tx(s,x,3.18,4.33,.31,n,15,TEAL,True);tx(s,x,3.86,4.31,.43,title,25,INK,True);tx(s,x,4.63,4.31,.76,sub,18,GRAY)
    ln(s,x,5.74,x+4.28,5.74,LINE,1);tx(s,x,6.08,4.31,.86,result,23,TEAL,True)
box(s,.85,7.26,14.23,.51,PALE);tx(s,1.06,7.38,13.81,.29,'실행 가능한 대안이 없으면  ·  대체 조달 / 기한 협의 / 예산 예외를 담당자가 결정',17,TEAL,True)
foot(s,'코드에서 경계값·비용 한도·지연·재고·적재·규격·예산 부족·동점·잘못된 입력을 확인. 프로토타입 검증이며 현장 운용 안정성 검증은 아닙니다.')
note(s,'비용 한도120으로 C가 제외되면 A와 B끼리 재정규화한다. B=(0,1,.5), Carbon RPS=.7이다. 원래 세 후보에서 B의 점수.67과 구분한다. 실제 배출오차 때문에 탄소예산 준수는 보장하지 않는다.')

# APPENDIX 4 / the serious limitation is visible, with a scoped extension.
s=base(p,4,'APPENDIX / LIMITS','월초 50%와 월말 50%는 같은 상태가 아닙니다.',
       '기본 CUR는 누적 소진량만 표현합니다. 남은 기간과 앞으로의 요청을 직접 설명하지 못합니다.',appendix=True)
for y,label,pos in [(3.63,'기간 초반',.18),(5.20,'기간 후반',.85)]:
    tx(s,.93,y,2.41,.39,label,23,INK,True);ln(s,3.80,y+.22,10.27,y+.22,LINE,5);circle(s,3.8+6.47*pos-.10,y+.12,.2,TEAL)
    tx(s,11.04,y-.09,4.05,.61,'CUR 50%',33,TEAL,True)
tx(s,.95,6.81,14.07,.78,'확장 방향: 남은 수요를 고려한 기말 예상 배출 점검\n기본 모형과 분리해 성능·예측 오차를 검증합니다.',23,INK,True)
foot(s,'추가 한계: 구간 경계에서 가중치 급변, 후보 집합 변화에 따른 상대 점수 변화, 배출 추정 오차. 전역 최적해나 기간 예산 준수를 보장하지 않습니다.')
note(s,'이 한계를 숨기지 않는다. 기말 예상 배출 기반 보조경보 또는 수요 대비 예산 사용 속도 지표를 검토하되, 데이터 없는 정교한 예측을 이미 구현했다고 주장하지 않는다. 임계값·가중치 민감도는 실제 요청 재생에서 평가한다.')

# APPENDIX 5 / a credible, fair evaluation protocol.
s=base(p,5,'APPENDIX / EVALUATION','같은 출발점에서, 정책별로 끝까지 비교해야 합니다.',
       '한 번의 가상 추천 결과는 기간 전체의 운영 효과를 증명하지 않습니다.',appendix=True)
tx(s,.92,3.15,4.24,.39,'동일하게 제공',23,TEAL,True)
tx(s,.94,3.97,4.13,2.12,'초기 재고 · 탄소예산\n요청 순서 · 후보 생성 규칙\n사용 가능한 정보\n수량·기한·비용·탄소 제약',20,INK)
arrow(s,5.29,4.86,.61)
for y,t in [(3.72,'비용 우선'),(4.66,'고정 가중치'),(5.60,'CBP')]:
    box(s,6.17,y,3.35,.66,PALE if t=='CBP' else WHITE);tx(s,6.41,y+.15,2.87,.37,t,23,TEAL if t=='CBP' else INK,True)
arrow(s,9.84,4.86,.61)
tx(s,10.83,3.15,4.35,.39,'기간 전체로 평가',23,TEAL,True)
tx(s,10.85,3.97,4.27,2.12,'총비용 · 실제 배출\n정시 공급 충족률\n기간말 탄소예산 초과\n예외 발생 횟수',20,INK)
box(s,.92,7.05,14.22,.62,NIGHT);tx(s,1.17,7.22,13.72,.34,'각 정책의 선택 이후에는 재고·누적 배출을 따로 갱신합니다. 결과에 유리한 미래 정보를 미리 주지 않습니다.',16,WHITE,True)
foot(s,'검증 순서: 과거 요청 재생 → 실제 배차에 영향을 주지 않는 추천 → 제한적 현장 실행. 비용·탄소·공급 중 불리해지는 항목도 함께 보고합니다.')
note(s,'고정가중치는0.5,0.3,0.2를 하나의 대조군으로 두되, 다른 고정가중치와도 비교해야 한다. 파라미터 튜닝 자료와 최종 평가 자료를 분리하여 과적합을 줄인다. 가상 후보를 실제 견적·운행·배출자료로 대체하기 전 성과율을 발표하지 않는다.')

# APPENDIX 6 / evidence provenance, with no invented literature.
s=base(p,6,'APPENDIX / EVIDENCE','제안한 구조와, 아직 입증하지 않은 주장을 구분합니다.',
       '가중합 자체의 발명을 주장하지 않습니다. ULS 재배치에서 탄소예산 상태를 선택 규칙에 연결하는 제안입니다.',appendix=True)
for x,title,items,c in [(.91,'현재 확인한 근거','예선 원안: CUR · 4모드 · RPS\n공식 안내: 4분 발표 · 배너 4쪽 발췌\n재현 코드: 점수·제약·예외 추천\n동작 데모: 사용률·비용 한도 변경',TEAL),(8.38,'추가 확인이 필요한 근거','현장 탄소 감축률 · 비용 효과\n동적 가중치·탄소예산 관련 선행 연구\n기존 풀링 시스템 대비 신규성\n역대 수상작의 심사·수상 근거',GRAY)]:
    tx(s,x,3.21,6.59,.40,title,25,c,True);ln(s,x,3.91,x+6.49,3.91,LINE,1);tx(s,x,4.28,6.54,2.71,items,21,INK)
box(s,.91,7.22,14.03,.54,PALE);tx(s,1.15,7.35,13.55,.29,'현재 단계의 주장: “작동하는 상태 연동 추천 규칙을 제안하고 구현했다.”',18,TEAL,True)
foot(s,'원자료: 두런두런 예선 PPT, KPCA 2026-1007 합격 안내, 2025 배너 예시. 외부 원문 접근 차단으로 논문·역대 수상작 비교는 완료하지 못했습니다.')
note(s,'외부 연구 접근은 Crossref/OpenAlex/Semantic Scholar/arXiv/KPCA/검색사이트에서 403 차단을 확인했다. 확인하지 않은 연구를 관련연구 표로 만들거나 신규성 입증이라고 말하지 않는다. 연구 접근에 필요한 환경설정 초안은 별도로 저장했다.')

save(p,OUT/'cbp-full-v4.pptx')
excerpt(p,list(range(5,11)),OUT/'cbp-research-v4.pptx')
text='# CBP V4 · 4분 발표 대본\n\n목표 3분 50초 + 전환·호흡 여유 10초. 실제 낭독 시간은 발표자가 재서 조정해야 합니다. 본문 5장만 발표하고 보충자료는 질의응답에 사용하세요.\n\n'
for t,title,script in SCRIPTS:text+='## '+t+' | '+title+'\n\n'+script+'\n\n'
text+='## 전달할 때\n\n- 1장은 제목을 읽은 뒤 바로 운영 질문으로 넘어갑니다.\n- 2장은 상태→기준→후보→다음 요청 순서만 손으로 짚습니다.\n- 3장에서는 세 후보 수치를 전부 반복하지 말고 C를 고르는 이유와 비용 한도를 강조합니다.\n- 4장은 CBP 행의 A→B→C만 따라가면 됩니다. 이를 실제 산업 전체와의 우위 비교로 말하지 않습니다.\n- 5장은 구현과 현장 효과 검증을 분리하고 마지막 문장 뒤에 멈춥니다.\n'
(OUT/'발표대본_4분.md').write_text(text)
print('Built main 5 / exact banner 4 / research 6 / combined 11; model checks:',verify())
